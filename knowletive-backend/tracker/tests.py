import io
import json
import zoneinfo
import openpyxl
from datetime import datetime, timezone, timedelta
from django.test import TestCase
from django.urls import reverse
from django.core.management import call_command
from rest_framework.test import APIClient
from rest_framework.authtoken.models import Token
from .models import User, Invitation, Notification
from .storage_adapter import storage

IST_TZ = zoneinfo.ZoneInfo("Asia/Kolkata")


class TimeTrackerTestCase(TestCase):
    def setUp(self):
        self.client = APIClient()
        
        # Create Admin
        self.admin = User.objects.create_superuser(
            username='admin@knowletive.com',
            email='admin@knowletive.com',
            password='adminpassword',
            first_name='Admin',
            last_name='User',
            role='ADMIN'
        )
        self.admin_token, _ = Token.objects.get_or_create(user=self.admin)

        # Create Employee
        self.employee = User.objects.create_user(
            username='emp@knowletive.com',
            email='emp@knowletive.com',
            password='emppassword',
            first_name='John',
            last_name='Doe',
            role='EMPLOYEE',
            department='Engineering',
            employee_id='JOHN001'
        )
        self.employee_token, _ = Token.objects.get_or_create(user=self.employee)

        # Clean mock file for isolated tests
        storage._ensure_mock_file()
        if storage.use_mock:
            import json
            with open(storage.local_file, 'w', encoding='utf-8') as f:
                json.dump({'TimeEvents': [], 'AuditLog': [], 'Settings': {'work_start_time': '10:00', 'work_end_time': '19:00', 'timezone': 'Asia/Kolkata'}}, f)

    def test_auth_login_and_logout(self):
        # Invalid credentials
        res = self.client.post('/api/auth/login/', {'email': 'emp@knowletive.com', 'password': 'wrong'})
        self.assertEqual(res.status_code, 400)

        # Valid login
        res = self.client.post('/api/auth/login/', {'email': 'emp@knowletive.com', 'password': 'emppassword'})
        self.assertEqual(res.status_code, 200)
        self.assertIn('token', res.data)
        token = res.data['token']

        # Logout
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {token}')
        res = self.client.post('/api/auth/logout/')
        self.assertEqual(res.status_code, 200)

    def test_permission_enforcement(self):
        # Unauthenticated request to tracker
        res = self.client.get('/api/tracker/status/')
        self.assertEqual(res.status_code, 401)

        # Employee cannot access admin dashboard
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        res = self.client.get('/api/admin/dashboard/')
        self.assertEqual(res.status_code, 403)

        # Admin can access admin dashboard
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        res = self.client.get('/api/admin/dashboard/')
        self.assertEqual(res.status_code, 200)

    def test_time_tracking_valid_sequence(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        
        # 1. Initial status: Not Started
        res = self.client.get('/api/tracker/status/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['status'], 'Not Started')

        # 2. Start Work
        res = self.client.post('/api/tracker/action/', {'action': 'START_WORK', 'source': 'web'})
        self.assertEqual(res.status_code, 200)

        res = self.client.get('/api/tracker/status/')
        self.assertEqual(res.data['status'], 'Working')
        self.assertIsNotNone(res.data['start_time'])

        # 3. Start Break
        res = self.client.post('/api/tracker/action/', {'action': 'START_BREAK', 'source': 'web'})
        self.assertEqual(res.status_code, 200)

        res = self.client.get('/api/tracker/status/')
        self.assertEqual(res.data['status'], 'On Break')

        # 4. Resume Work
        res = self.client.post('/api/tracker/action/', {'action': 'RESUME_WORK', 'source': 'web'})
        self.assertEqual(res.status_code, 200)

        res = self.client.get('/api/tracker/status/')
        self.assertEqual(res.data['status'], 'Working')

        # 5. End Work
        res = self.client.post('/api/tracker/action/', {'action': 'END_WORK', 'source': 'web'})
        self.assertEqual(res.status_code, 200)

        res = self.client.get('/api/tracker/status/')
        self.assertEqual(res.data['status'], 'Ended')
        self.assertIsNotNone(res.data['end_time'])

    def test_invalid_state_transitions(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')

        # Cannot start break before starting work
        res = self.client.post('/api/tracker/action/', {'action': 'START_BREAK'})
        self.assertEqual(res.status_code, 400)

        # Cannot resume work before starting work
        res = self.client.post('/api/tracker/action/', {'action': 'RESUME_WORK'})
        self.assertEqual(res.status_code, 400)

        # Cannot end work before starting work
        res = self.client.post('/api/tracker/action/', {'action': 'END_WORK'})
        self.assertEqual(res.status_code, 400)

        # Start work
        self.client.post('/api/tracker/action/', {'action': 'START_WORK'})

        # Cannot start work again while already working
        res = self.client.post('/api/tracker/action/', {'action': 'START_WORK'})
        self.assertEqual(res.status_code, 400)

    def test_admin_correction_workflow_and_audit(self):
        # Employee records START_WORK
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        act_res = self.client.post('/api/tracker/action/', {'action': 'START_WORK'})
        event_id = act_res.data['event']['event_id']

        # Switch to Admin
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')

        # Attempt correction without mandatory reason
        res = self.client.post('/api/admin/corrections/', {
            'event_id': event_id,
            'new_timestamp': '2026-10-07T10:00:00Z',
            'reason': ''
        })
        self.assertEqual(res.status_code, 400)
        self.assertIn('strictly required', res.data['error'])

        # Valid correction with reason
        res = self.client.post('/api/admin/corrections/', {
            'event_id': event_id,
            'new_timestamp': '2026-10-07T10:00:00Z',
            'reason': 'Adjusting for missed system clock-in due to power outage'
        })
        self.assertEqual(res.status_code, 200)

        # Verify audit log recorded the correction
        res = self.client.get('/api/admin/audit-logs/')
        self.assertEqual(res.status_code, 200)
        logs = res.data['audit_logs']
        correction_log = next((l for l in logs if l['action'] == 'ADMIN_CORRECTION'), None)
        self.assertIsNotNone(correction_log)
        self.assertIn('Adjusting for missed system clock-in', correction_log['details'])

    def test_reports_and_export_formats(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')

        # Fetch JSON reports
        res = self.client.get('/api/admin/reports/')
        self.assertEqual(res.status_code, 200)
        self.assertIn('reports', res.data)

        # CSV Export
        res = self.client.get('/api/admin/reports/export/?format=csv')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers['Content-Type'], 'text/csv; charset=utf-8')
        csv_text = res.content.decode('utf-8')
        self.assertIn('Date (IST)', csv_text)
        self.assertIn('Break Duration', csv_text)
        self.assertIn('Break Start/End Details', csv_text)

        # Excel Export
        res = self.client.get('/api/admin/reports/export/?format=xlsx')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertGreater(len(res.content), 100)

        # Verify Excel headers
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        headers = [cell.value for cell in ws[4]]
        self.assertIn('Break Duration', headers)
        self.assertIn('Break Start/End Details', headers)

    def test_export_break_duration_columns_and_dynamic_filenames(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')

        # Create two distinct test employees
        emp1 = User.objects.create_user(
            username='ayush.kadam@knowletive.com',
            email='ayush.kadam@knowletive.com',
            password='Password123!',
            first_name='Ayush',
            last_name='Kadam',
            role='EMPLOYEE',
            department='Engineering',
            employee_id='EMP101'
        )
        emp2 = User.objects.create_user(
            username='aman.bendkule@knowletive.com',
            email='aman.bendkule@knowletive.com',
            password='Password123!',
            first_name='Aman',
            last_name='Bendkule',
            role='EMPLOYEE',
            department='Product',
            employee_id='EMP102'
        )

        # Populate recorded time events with breaks
        # Emp 1 (Ayush Kadam) on 2026-10-05:
        # 09:00 IST (03:30 UTC) - START_WORK
        # 13:00 IST (07:30 UTC) - START_BREAK
        # 13:45 IST (08:15 UTC) - RESUME_WORK  (45 min break)
        # 16:00 IST (10:30 UTC) - START_BREAK
        # 16:15 IST (10:45 UTC) - RESUME_WORK  (15 min break)
        # 18:00 IST (12:30 UTC) - END_WORK     (Total break = 1 hr = 01:00:00)
        # Emp 1 on 2026-10-06:
        # 09:30 IST (04:00 UTC) - START_WORK
        # 13:30 IST (08:00 UTC) - START_BREAK
        # 14:00 IST (08:30 UTC) - RESUME_WORK  (30 min break)
        # 18:30 IST (13:00 UTC) - END_WORK     (Total break = 30 min = 00:30:00)

        # Emp 2 (Aman Bendkule) on 2026-10-05:
        # 10:00 IST (04:30 UTC) - START_WORK
        # 12:30 IST (07:00 UTC) - START_BREAK
        # 13:00 IST (07:30 UTC) - RESUME_WORK  (30 min break)
        # 19:00 IST (13:30 UTC) - END_WORK     (Total break = 30 min = 00:30:00)

        events_data = [
            # Ayush - Oct 5
            {'event_id': 'e1', 'user_id': emp1.id, 'event_type': 'START_WORK', 'timestamp': '2026-10-05T03:30:00Z', 'source': 'web'},
            {'event_id': 'e2', 'user_id': emp1.id, 'event_type': 'START_BREAK', 'timestamp': '2026-10-05T07:30:00Z', 'source': 'web'},
            {'event_id': 'e3', 'user_id': emp1.id, 'event_type': 'RESUME_WORK', 'timestamp': '2026-10-05T08:15:00Z', 'source': 'web'},
            {'event_id': 'e4', 'user_id': emp1.id, 'event_type': 'START_BREAK', 'timestamp': '2026-10-05T10:30:00Z', 'source': 'web'},
            {'event_id': 'e5', 'user_id': emp1.id, 'event_type': 'RESUME_WORK', 'timestamp': '2026-10-05T10:45:00Z', 'source': 'web'},
            {'event_id': 'e6', 'user_id': emp1.id, 'event_type': 'END_WORK', 'timestamp': '2026-10-05T12:30:00Z', 'source': 'web'},
            # Ayush - Oct 6
            {'event_id': 'e7', 'user_id': emp1.id, 'event_type': 'START_WORK', 'timestamp': '2026-10-06T04:00:00Z', 'source': 'web'},
            {'event_id': 'e8', 'user_id': emp1.id, 'event_type': 'START_BREAK', 'timestamp': '2026-10-06T08:00:00Z', 'source': 'web'},
            {'event_id': 'e9', 'user_id': emp1.id, 'event_type': 'RESUME_WORK', 'timestamp': '2026-10-06T08:30:00Z', 'source': 'web'},
            {'event_id': 'e10', 'user_id': emp1.id, 'event_type': 'END_WORK', 'timestamp': '2026-10-06T13:00:00Z', 'source': 'web'},
            # Aman - Oct 5
            {'event_id': 'e11', 'user_id': emp2.id, 'event_type': 'START_WORK', 'timestamp': '2026-10-05T04:30:00Z', 'source': 'web'},
            {'event_id': 'e12', 'user_id': emp2.id, 'event_type': 'START_BREAK', 'timestamp': '2026-10-05T07:00:00Z', 'source': 'web'},
            {'event_id': 'e13', 'user_id': emp2.id, 'event_type': 'RESUME_WORK', 'timestamp': '2026-10-05T07:30:00Z', 'source': 'web'},
            {'event_id': 'e14', 'user_id': emp2.id, 'event_type': 'END_WORK', 'timestamp': '2026-10-05T13:30:00Z', 'source': 'web'},
        ]

        # Save to mock storage
        import json
        with open(storage.local_file, 'r+', encoding='utf-8') as f:
            data = json.load(f)
            data['TimeEvents'] = events_data
            f.seek(0)
            json.dump(data, f, indent=2)
            f.truncate()

        # 1. TEST ADMIN PANEL JSON REPORTS API
        rep_res = self.client.get(f'/api/admin/reports/?user_id={emp1.id}&start_date=2026-10-01&end_date=2026-10-07')
        self.assertEqual(rep_res.status_code, 200)
        self.assertEqual(len(rep_res.data['reports']), 7)
        r_oct5 = next(r for r in rep_res.data['reports'] if r['date'] == '2026-10-05')
        self.assertEqual(r_oct5['employee_name'], 'Ayush Kadam')
        self.assertEqual(r_oct5['break_duration'], '01:00:00')
        self.assertIn('01:00:00 PM - 01:45:00 PM', r_oct5['break_details'])
        self.assertIn('04:00:00 PM - 04:15:00 PM', r_oct5['break_details'])
        r_oct6 = next(r for r in rep_res.data['reports'] if r['date'] == '2026-10-06')
        self.assertEqual(r_oct6['break_duration'], '00:30:00')
        self.assertEqual(rep_res.data['summary']['total_break_duration'], '01:30:00')

        # 2. TEST AYUSH KADAM CSV EXPORT (Single Employee, Date Range)
        csv_emp1 = self.client.get(f'/api/admin/reports/export/?format=csv&user_id={emp1.id}&start_date=2026-10-01&end_date=2026-10-07')
        self.assertEqual(csv_emp1.status_code, 200)
        self.assertIn('filename="Ayush_Kadam_Time_Report_2026-10-01_to_2026-10-07.csv"', csv_emp1.headers['Content-Disposition'])
        self.assertEqual(csv_emp1.headers['X-Suggested-Filename'], 'Ayush_Kadam_Time_Report_2026-10-01_to_2026-10-07.csv')
        csv1_content = csv_emp1.content.decode('utf-8')
        self.assertIn('Ayush Kadam', csv1_content)
        self.assertNotIn('Aman Bendkule', csv1_content)  # Filter correctly excludes other employee
        self.assertIn('01:00:00', csv1_content)
        self.assertIn('01:00:00 PM - 01:45:00 PM', csv1_content)
        self.assertIn('01:30:00', csv1_content)  # Totals row break duration

        # 3. TEST AYUSH KADAM EXCEL EXPORT (Single Employee, Date Range)
        xlsx_emp1 = self.client.get(f'/api/admin/reports/export/?format=xlsx&user_id={emp1.id}&start_date=2026-10-01&end_date=2026-10-07')
        self.assertEqual(xlsx_emp1.status_code, 200)
        self.assertIn('filename="Ayush_Kadam_Time_Report_2026-10-01_to_2026-10-07.xlsx"', xlsx_emp1.headers['Content-Disposition'])
        self.assertEqual(xlsx_emp1.headers['X-Suggested-Filename'], 'Ayush_Kadam_Time_Report_2026-10-01_to_2026-10-07.xlsx')
        wb1 = openpyxl.load_workbook(io.BytesIO(xlsx_emp1.content))
        ws1 = wb1.active
        headers1 = [cell.value for cell in ws1[4]]
        self.assertIn('Break Duration', headers1)
        self.assertIn('Break Start/End Details', headers1)
        date_idx1 = headers1.index('Date (IST)') + 1
        break_dur_idx = headers1.index('Break Duration') + 1
        break_det_idx = headers1.index('Break Start/End Details') + 1
        emp_name_idx = headers1.index('Employee Name') + 1
        
        # Check Oct 5 and Oct 6 rows in Excel
        row_oct5 = next(r for r in range(5, ws1.max_row) if ws1.cell(row=r, column=date_idx1).value == '2026-10-05')
        self.assertEqual(ws1.cell(row=row_oct5, column=emp_name_idx).value, 'Ayush Kadam')
        self.assertEqual(ws1.cell(row=row_oct5, column=break_dur_idx).value, '01:00:00')
        self.assertIn('01:00:00 PM - 01:45:00 PM', ws1.cell(row=row_oct5, column=break_det_idx).value)

        row_oct6 = next(r for r in range(5, ws1.max_row) if ws1.cell(row=r, column=date_idx1).value == '2026-10-06')
        self.assertEqual(ws1.cell(row=row_oct6, column=break_dur_idx).value, '00:30:00')

        # 4. TEST AMAN BENDKULE EXCEL & CSV EXPORT (Single Employee, Full Month)
        xlsx_emp2 = self.client.get(f'/api/admin/reports/export/?format=xlsx&user_id={emp2.id}&start_date=2026-10-01&end_date=2026-10-31')
        self.assertEqual(xlsx_emp2.status_code, 200)
        self.assertIn('filename="Aman_Bendkule_Time_Report_2026-10-01_to_2026-10-31.xlsx"', xlsx_emp2.headers['Content-Disposition'])
        wb2 = openpyxl.load_workbook(io.BytesIO(xlsx_emp2.content))
        ws2 = wb2.active
        headers2 = [cell.value for cell in ws2[4]]
        date_idx2 = headers2.index('Date (IST)') + 1
        row_oct5_aman = next(r for r in range(5, ws2.max_row) if ws2.cell(row=r, column=date_idx2).value == '2026-10-05')
        self.assertEqual(ws2.cell(row=row_oct5_aman, column=emp_name_idx).value, 'Aman Bendkule')
        self.assertEqual(ws2.cell(row=row_oct5_aman, column=break_dur_idx).value, '00:30:00')
        self.assertEqual(ws2.cell(row=row_oct5_aman, column=break_det_idx).value, '12:30:00 PM - 01:00:00 PM')

        csv_emp2 = self.client.get(f'/api/admin/reports/export/?format=csv&user_id={emp2.id}&start_date=2026-10-01&end_date=2026-10-31')
        self.assertEqual(csv_emp2.status_code, 200)
        self.assertIn('filename="Aman_Bendkule_Time_Report_2026-10-01_to_2026-10-31.csv"', csv_emp2.headers['Content-Disposition'])
        csv2_content = csv_emp2.content.decode('utf-8')
        self.assertIn('Aman Bendkule', csv2_content)
        self.assertNotIn('Ayush Kadam', csv2_content)

        # 5. TEST ALL EMPLOYEES EXPORT (Date Range)
        csv_all = self.client.get('/api/admin/reports/export/?format=csv&start_date=2026-10-01&end_date=2026-10-07')
        self.assertEqual(csv_all.status_code, 200)
        self.assertIn('filename="All_Employees_Time_Report_2026-10-01_to_2026-10-07.csv"', csv_all.headers['Content-Disposition'])
        csv_all_content = csv_all.content.decode('utf-8')
        self.assertIn('Ayush Kadam', csv_all_content)
        self.assertIn('Aman Bendkule', csv_all_content)

        xlsx_all = self.client.get('/api/admin/reports/export/?format=xlsx&start_date=2026-10-01&end_date=2026-10-07')
        self.assertEqual(xlsx_all.status_code, 200)
        self.assertIn('filename="All_Employees_Time_Report_2026-10-01_to_2026-10-07.xlsx"', xlsx_all.headers['Content-Disposition'])

        # 6. TEST MONTH FILTER MODE
        csv_month = self.client.get(f'/api/admin/reports/export/?format=csv&user_id={emp1.id}&date_mode=month&month=10&year=2026')
        self.assertEqual(csv_month.status_code, 200)
        self.assertIn('filename="Ayush_Kadam_Time_Report_2026-10-01_to_2026-10-31.csv"', csv_month.headers['Content-Disposition'])

        # 7. TEST DATE FILTERING EXCLUDES OUT-OF-RANGE DATES
        csv_range_filtered = self.client.get('/api/admin/reports/export/?format=csv&start_date=2026-10-06&end_date=2026-10-07')
        self.assertEqual(csv_range_filtered.status_code, 200)
        csv_filtered_content = csv_range_filtered.content.decode('utf-8')
        self.assertIn('2026-10-06', csv_filtered_content)
        self.assertNotIn('2026-10-05', csv_filtered_content)

    def test_user_management(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')

        # Create new employee
        res = self.client.post('/api/admin/users/', {
            'email': 'alice@knowletive.com',
            'first_name': 'Alice',
            'last_name': 'Smith',
            'department': 'Design',
            'role': 'EMPLOYEE',
            'password': 'SecurePassword123!'
        })
        self.assertEqual(res.status_code, 201)
        new_user_id = res.data['user']['id']

        # Edit employee
        res = self.client.patch(f'/api/admin/users/{new_user_id}/', {
            'department': 'Product',
            'employment_status': 'Part-Time'
        })
        self.assertEqual(res.status_code, 200)

        # Deactivate employee
        res = self.client.patch(f'/api/admin/users/{new_user_id}/', {
            'is_active': False
        })
        self.assertEqual(res.status_code, 200)

        # Deactivated user cannot log in
        client2 = APIClient()
        res = client2.post('/api/auth/login/', {'email': 'alice@knowletive.com', 'password': 'SecurePassword123!'})
        self.assertEqual(res.status_code, 403)

    def test_monthly_timesheet_permissions_and_entries(self):
        # 1. Employee fetches own timesheet
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        res = self.client.get('/api/timesheet/monthly/?year=2026&month=10')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['employee']['id'], self.employee.id)
        self.assertIn('days', res.data)
        self.assertIn('summary', res.data)

        # 2. Employee tries to access Admin's timesheet -> 403 Forbidden
        res = self.client.get(f'/api/timesheet/monthly/?year=2026&month=10&user_id={self.admin.id}')
        self.assertEqual(res.status_code, 403)
        self.assertIn('Access denied', res.data['error'])

        # 3. Employee updates project, task, description for a work date
        res = self.client.post('/api/timesheet/entry/', {
            'date': '2026-10-07',
            'project_name': 'Knowletive Website',
            'task_name': 'Website Redesign',
            'work_description': 'Implemented responsive monthly timesheet grid'
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['entry']['project_name'], 'Knowletive Website')

        # 4. Employee tries to update another user's timesheet entry -> 403 Forbidden
        res = self.client.post('/api/timesheet/entry/', {
            'user_id': self.admin.id,
            'date': '2026-10-07',
            'project_name': 'Hacked Project',
            'task_name': 'Unauthorized',
            'work_description': 'Malicious payload'
        })
        self.assertEqual(res.status_code, 403)

        # 5. Admin can access employee's timesheet and review entries
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        res = self.client.get(f'/api/timesheet/monthly/?year=2026&month=10&user_id={self.employee.id}')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['employee']['id'], self.employee.id)
        
        # Verify the saved entry is present
        target_day = next((d for d in res.data['days'] if d['date'] == '2026-10-07'), None)
        self.assertIsNotNone(target_day)
        self.assertEqual(target_day['project_name'], 'Knowletive Website')
        self.assertEqual(target_day['task_name'], 'Website Redesign')

    def test_timesheet_exports(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        
        # Export CSV
        res = self.client.get('/api/timesheet/export/?year=2026&month=10&format=csv')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers['Content-Type'], 'text/csv; charset=utf-8')
        content = res.content.decode('utf-8')
        self.assertIn('KNOWLETIVE MONTHLY EMPLOYEE TIMESHEET', content)
        self.assertIn('Sr. No.', content)

        # Export Excel (.xlsx)
        res = self.client.get('/api/timesheet/export/?year=2026&month=10&format=xlsx')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertGreater(len(res.content), 200)

    def test_timewise_task_logging_and_admin_reports(self):
        # 1. Employee starts first task
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        res = self.client.post('/api/tasks/start/', {
            'project_name': 'Project Alpha',
            'task_name': 'Reviewing customer data',
            'description': 'Analyzing CSV imports'
        })
        self.assertEqual(res.status_code, 201)
        task1 = res.data['task']
        self.assertEqual(task1['project_name'], 'Project Alpha')
        self.assertEqual(task1['status'], 'IN_PROGRESS')
        self.assertIsNotNone(task1['start_time'])
        task1_id = task1['task_log_id']

        # 2. Cannot start another task while one is currently in progress
        res = self.client.post('/api/tasks/start/', {
            'project_name': 'Project Beta',
            'task_name': 'Preparing report',
            'description': 'Drafting Q3 deck'
        })
        self.assertEqual(res.status_code, 400)
        self.assertIn('already in progress', res.data['error'])

        # 3. Employee updates project/task/description of current task
        res = self.client.patch(f'/api/tasks/{task1_id}/', {
            'description': 'Updated description: Analyzing customer churn metrics'
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['task']['description'], 'Updated description: Analyzing customer churn metrics')

        # 4. End first task
        res = self.client.post('/api/tasks/end/', {
            'task_log_id': task1_id
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['task']['status'], 'COMPLETED')
        self.assertIsNotNone(res.data['task']['end_time'])

        # 5. Start second task on the same day
        res = self.client.post('/api/tasks/start/', {
            'project_name': 'Project Beta',
            'task_name': 'Preparing report',
            'description': 'Drafting Q3 deck'
        })
        self.assertEqual(res.status_code, 201)
        task2_id = res.data['task']['task_log_id']

        # 6. Fetch daily tasks for employee
        res = self.client.get('/api/tasks/daily/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data['tasks']), 2)
        self.assertEqual(res.data['tasks'][0]['project_name'], 'Project Alpha')
        self.assertEqual(res.data['tasks'][1]['project_name'], 'Project Beta')
        self.assertIn('project_breakdown', res.data)
        self.assertGreaterEqual(len(res.data['project_breakdown']), 2)

        # 7. Employee forbidden from viewing other user's daily tasks
        res = self.client.get(f'/api/tasks/daily/?user_id={self.admin.id}')
        self.assertEqual(res.status_code, 403)

        # 8. Admin view: report with filters
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        res = self.client.get('/api/admin/tasks/report/')
        self.assertEqual(res.status_code, 200)
        self.assertIn('tasks', res.data)
        self.assertIn('by_employee', res.data)
        self.assertIn('by_project_and_employee', res.data)
        self.assertIn('summary', res.data)

        # Filter by project
        res = self.client.get('/api/admin/tasks/report/?project=Project%20Alpha')
        self.assertEqual(res.status_code, 200)
        self.assertTrue(all(t['project_name'] == 'Project Alpha' for t in res.data['tasks']))

        # 9. Admin CSV export
        res = self.client.get('/api/admin/tasks/export/?format=csv')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers['Content-Type'], 'text/csv; charset=utf-8')
        content = res.content.decode('utf-8')
        self.assertIn('KNOWLETIVE DAILY TIME-WISE TASK REPORT', content)
        self.assertIn('Project Alpha', content)

        # 10. Admin Excel export
        res = self.client.get('/api/admin/tasks/export/?format=xlsx')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertGreater(len(res.content), 200)


class EmployeeSignupAndOnboardingTestCase(TestCase):
    def setUp(self):
        self.client = APIClient()

        # Create Admin
        self.admin = User.objects.create_superuser(
            username='admin@knowletive.com',
            email='admin@knowletive.com',
            password='AdminPassword123!',
            first_name='Admin',
            last_name='User',
            role='ADMIN',
            approval_status='APPROVED'
        )
        self.admin_token, _ = Token.objects.get_or_create(user=self.admin)

        # Create existing Employee
        self.employee = User.objects.create_user(
            username='existing@knowletive.com',
            email='existing@knowletive.com',
            password='EmployeePassword123!',
            first_name='Existing',
            last_name='Staff',
            role='EMPLOYEE',
            approval_status='APPROVED'
        )
        self.employee_token, _ = Token.objects.get_or_create(user=self.employee)

        # Create test invitation
        self.invitation = Invitation.objects.create(
            code='KNOW-TEST1234',
            created_by=self.admin,
            expires_at=datetime.now(timezone.utc) + timedelta(days=7),
            notes='Test invitation'
        )

        storage._ensure_mock_file()
        if storage.use_mock:
            import json
            with open(storage.local_file, 'w', encoding='utf-8') as f:
                json.dump({'TimeEvents': [], 'TaskLogs': [], 'TimesheetEntries': [], 'AuditLog': [], 'Settings': {'work_start_time': '10:00', 'work_end_time': '19:00', 'timezone': 'Asia/Kolkata'}}, f)

    def test_admin_generate_invitations(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        
        # 1. Admin generates an invitation
        res = self.client.post('/api/admin/invitations/', {
            'recipient_email': 'newhire@knowletive.com',
            'days_valid': 14,
            'notes': 'Engineering candidate'
        })
        self.assertEqual(res.status_code, 201)
        self.assertIn('code', res.data['invitation'])
        code = res.data['invitation']['code']
        self.assertTrue(code.startswith('KNOW-'))

        # 2. Admin lists invitations
        res = self.client.get('/api/admin/invitations/')
        self.assertEqual(res.status_code, 200)
        self.assertGreaterEqual(len(res.data['invitations']), 2)

        # 3. Employee forbidden from generating or listing invitations
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        res = self.client.get('/api/admin/invitations/')
        self.assertEqual(res.status_code, 403)
        res = self.client.post('/api/admin/invitations/', {'days_valid': 7})
        self.assertEqual(res.status_code, 403)

    def test_public_signup_strictly_assigns_employee_role(self):
        # 1. Standard public signup receives EMPLOYEE role
        res1 = self.client.post('/api/auth/signup/', {
            'username': 'standard_worker',
            'password': 'password123'
        })
        self.assertEqual(res1.status_code, 201)
        self.assertEqual(res1.data['user']['role'], 'EMPLOYEE')

        worker = User.objects.get(username='standard_worker')
        self.assertEqual(worker.role, 'EMPLOYEE')
        self.assertFalse(worker.is_staff)
        self.assertFalse(worker.is_superuser)
        self.assertTrue(worker.check_password('password123'))
        self.assertNotEqual(worker.password, 'password123')  # Securely hashed

        # 2. Malicious user attempting to request ADMIN role via public signup
        res2 = self.client.post('/api/auth/signup/', {
            'username': 'fake_admin',
            'password': 'password123',
            'role': 'ADMIN',
            'account_type': 'ADMIN'
        })
        self.assertEqual(res2.status_code, 201)
        # Backend MUST strictly override and enforce EMPLOYEE
        self.assertEqual(res2.data['user']['role'], 'EMPLOYEE')

        fake_admin = User.objects.get(username='fake_admin')
        self.assertEqual(fake_admin.role, 'EMPLOYEE')
        self.assertFalse(fake_admin.is_staff)
        self.assertFalse(fake_admin.is_superuser)

    def test_newly_signed_up_employee_forbidden_from_admin_apis(self):
        # Signup new employee
        signup_res = self.client.post('/api/auth/signup/', {
            'username': 'new_recruit',
            'password': 'password123'
        })
        self.assertEqual(signup_res.status_code, 201)
        token = signup_res.data['token']

        self.client.credentials(HTTP_AUTHORIZATION=f'Token {token}')

        # 1. Admin Dashboard -> 403
        res = self.client.get('/api/admin/dashboard/')
        self.assertEqual(res.status_code, 403)
        self.assertNotIn('overview', res.data)

        # 2. User Management -> 403
        res = self.client.get('/api/admin/users/')
        self.assertEqual(res.status_code, 403)

        # 3. Reports & Exports -> 403
        res = self.client.get('/api/admin/reports/')
        self.assertEqual(res.status_code, 403)
        res = self.client.get('/api/admin/reports/export/')
        self.assertEqual(res.status_code, 403)

        # 4. Daily Task Reports -> 403
        res = self.client.get('/api/admin/tasks/report/')
        self.assertEqual(res.status_code, 403)
        res = self.client.get('/api/admin/tasks/export/')
        self.assertEqual(res.status_code, 403)

        # 5. Audit Logs -> 403
        res = self.client.get('/api/admin/audit-logs/')
        self.assertEqual(res.status_code, 403)

        # 6. Admin Attendance Corrections -> 403
        res = self.client.post('/api/admin/corrections/', {'event_id': 'xyz', 'reason': 'test'})
        self.assertEqual(res.status_code, 403)

    def test_employee_cannot_access_other_employees_records(self):
        # Create Employee A and Employee B
        emp_a = User.objects.create_user(username='emp_a', password='password123', role='EMPLOYEE')
        emp_b = User.objects.create_user(username='emp_b', password='password123', role='EMPLOYEE')

        token_a, _ = Token.objects.get_or_create(user=emp_a)
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {token_a.key}')

        # 1. Employee A requests Employee B's monthly timesheet -> 403 Forbidden
        res_view = self.client.get(f'/api/timesheet/monthly/?user_id={emp_b.id}')
        self.assertEqual(res_view.status_code, 403)

        # 2. Employee A attempts to save timesheet entry for Employee B -> 403 Forbidden
        res_save = self.client.post('/api/timesheet/entry/', {
            'user_id': emp_b.id,
            'date': '2026-10-07',
            'project_name': 'Hacked',
            'task_name': 'Injected'
        })
        self.assertEqual(res_save.status_code, 403)

        # 3. Employee A attempts to export Employee B's timesheet -> 403 Forbidden
        res_export = self.client.get(f'/api/timesheet/export/?user_id={emp_b.id}')
        self.assertEqual(res_export.status_code, 403)

        # 4. Employee A can access their own timesheet -> 200 OK
        res_own = self.client.get(f'/api/timesheet/monthly/?user_id={emp_a.id}')
        self.assertEqual(res_own.status_code, 200)

    def test_first_admin_created_via_management_command(self):
        # Create first admin via setup_admin command
        call_command(
            'setup_admin',
            username='cli_admin',
            password='AdminPassword123!',
            no_input=True
        )

        admin_user = User.objects.get(username='cli_admin')
        self.assertEqual(admin_user.role, 'ADMIN')
        self.assertTrue(admin_user.is_staff)
        self.assertTrue(admin_user.is_superuser)
        self.assertTrue(admin_user.check_password('AdminPassword123!'))

        # Log in through login API
        res = self.client.post('/api/auth/login/', {
            'username': 'cli_admin',
            'password': 'AdminPassword123!'
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['user']['role'], 'ADMIN')

        # Admin CAN access admin dashboard
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {res.data["token"]}')
        dash_res = self.client.get('/api/admin/dashboard/')
        self.assertEqual(dash_res.status_code, 200)

        # Only existing Admin can create another Admin
        create_admin_res = self.client.post('/api/admin/users/', {
            'username': 'second_admin',
            'password': 'AdminPassword123!',
            'role': 'ADMIN'
        })
        self.assertEqual(create_admin_res.status_code, 201)
        self.assertEqual(create_admin_res.data['user']['role'], 'ADMIN')

    def test_duplicate_username_rejected(self):
        signup_data = {
            'username': 'unique_worker',
            'password': 'mypassword123'
        }
        res1 = self.client.post('/api/auth/signup/', signup_data)
        self.assertEqual(res1.status_code, 201)

        # Duplicate signup with same username (case-insensitive)
        signup_data_dup = {
            'username': 'UNIQUE_WORKER',
            'password': 'differentpassword'
        }
        res2 = self.client.post('/api/auth/signup/', signup_data_dup)
        self.assertEqual(res2.status_code, 400)
        self.assertIn('already taken', res2.data['error'].lower())

    def test_username_and_password_validation(self):
        # Empty username
        res = self.client.post('/api/auth/signup/', {
            'username': '',
            'password': 'validpassword123'
        })
        self.assertEqual(res.status_code, 400)
        self.assertIn('username is required', res.data['error'].lower())

        # Too short username
        res = self.client.post('/api/auth/signup/', {
            'username': 'ab',
            'password': 'validpassword123'
        })
        self.assertEqual(res.status_code, 400)
        self.assertIn('at least 3 characters', res.data['error'].lower())

        # Too short password
        res = self.client.post('/api/auth/signup/', {
            'username': 'validuser',
            'password': '123'
        })
        self.assertEqual(res.status_code, 400)
        self.assertIn('at least 6 characters', res.data['error'].lower())

    def test_login_with_username_and_wrong_password(self):
        User.objects.create_user(
            username='sammy',
            password='SecretPassword123!',
            role='EMPLOYEE'
        )

        # 1. Wrong password
        res_wrong = self.client.post('/api/auth/login/', {
            'username': 'sammy',
            'password': 'WrongPassword999'
        })
        self.assertEqual(res_wrong.status_code, 400)
        self.assertIn('invalid username or password', res_wrong.data['error'].lower())

        # 2. Correct password
        res_ok = self.client.post('/api/auth/login/', {
            'username': 'sammy',
            'password': 'SecretPassword123!'
        })
        self.assertEqual(res_ok.status_code, 200)
        self.assertIn('token', res_ok.data)
        self.assertEqual(res_ok.data['user']['role'], 'EMPLOYEE')
        self.assertEqual(res_ok.data['user']['username'], 'sammy')

    def test_deactivated_account_login_blocked(self):
        User.objects.create_user(
            username='disabled_guy',
            password='mypassword123',
            role='EMPLOYEE',
            is_active=False
        )
        res = self.client.post('/api/auth/login/', {
            'username': 'disabled_guy',
            'password': 'mypassword123'
        })
        self.assertEqual(res.status_code, 403)
        self.assertIn('deactivated', res.data['error'].lower())

    def test_admin_attendance_and_reports_exclude_admin_accounts(self):
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        
        # 1. Admin dashboard live attendance excludes admin accounts
        dash_res = self.client.get('/api/admin/dashboard/')
        self.assertEqual(dash_res.status_code, 200)
        
        # Total employees overview must exclude Admin
        total_employees = dash_res.data['overview']['total_employees']
        all_emp_count = User.objects.filter(role='EMPLOYEE').count()
        self.assertEqual(total_employees, all_emp_count)
        
        # Every user row must strictly have role == 'EMPLOYEE'
        returned_user_roles = [u['role'] for u in dash_res.data['users']]
        self.assertTrue(all(r == 'EMPLOYEE' for r in returned_user_roles))
        self.assertNotIn('ADMIN', returned_user_roles)

        # 2. Reports view strictly includes only employees
        rep_res = self.client.get('/api/admin/reports/')
        self.assertEqual(rep_res.status_code, 200)
        for r in rep_res.data['reports']:
            self.assertNotEqual(r['email'], self.admin.email)
            self.assertNotEqual(r['employee_id'], self.admin.id)

        # 3. CSV report export strictly contains only employees
        csv_res = self.client.get('/api/admin/reports/export/?format=csv')
        self.assertEqual(csv_res.status_code, 200)
        csv_content = csv_res.content.decode('utf-8')
        self.assertNotIn(self.admin.email, csv_content)

        # 4. Monthly timesheet for an Admin account is rejected (target must be an Employee)
        ts_res = self.client.get(f'/api/timesheet/monthly/?user_id={self.admin.id}')
        self.assertEqual(ts_res.status_code, 404)
        self.assertIn('Employee not found', ts_res.data['error'])

    def test_admin_tasks_report_strictly_shows_only_employee_tasks(self):
        # Record task for employee
        emp_task = storage.start_task(
            user_id=self.employee.id,
            project_name='Client Portal',
            task_name='Build Auth Flow',
            description='Employee work session',
            user_email=self.employee.email
        )
        storage.end_task(user_id=self.employee.id, task_log_id=emp_task['task_log_id'])

        # Record task for admin (simulate accidental admin task record in storage)
        admin_task = storage.start_task(
            user_id=self.admin.id,
            project_name='Admin Ops',
            task_name='System Configuration',
            description='Admin action',
            user_email=self.admin.email
        )
        storage.end_task(user_id=self.admin.id, task_log_id=admin_task['task_log_id'])

        # Admin fetches daily task report
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        res = self.client.get('/api/admin/tasks/report/')
        self.assertEqual(res.status_code, 200)
        
        # Verify tasks returned strictly exclude the admin's task
        returned_task_ids = [t['task_log_id'] for t in res.data['tasks']]
        self.assertIn(emp_task['task_log_id'], returned_task_ids)
        self.assertNotIn(admin_task['task_log_id'], returned_task_ids)

        # Verify summary counts only employee tasks
        self.assertEqual(res.data['summary']['total_tasks'], 1)
        self.assertEqual(len(res.data['by_employee']), 1)
        self.assertEqual(res.data['by_employee'][0]['user_id'], self.employee.id)
        self.assertEqual(res.data['by_employee'][0]['employee_id'], self.employee.employee_id or f"EMP{self.employee.id:03d}")

        # Verify CSV export excludes admin task
        csv_res = self.client.get('/api/admin/tasks/export/?format=csv')
        self.assertEqual(csv_res.status_code, 200)
        csv_text = csv_res.content.decode('utf-8')
        self.assertIn('Client Portal', csv_text)
        self.assertNotIn('Admin Ops', csv_text)
        self.assertIn('Employee ID', csv_text)

    def test_employee_id_generation_format_and_uniqueness(self):
        """
        Tests:
        - Uppercase name-based prefix + 3-digit number (e.g. AYUSH001, AMAN002, RAVINA003, AMAY004, SAKSHAM005).
        - Uniqueness: if two employees have the same name or prefix, use next available number.
        - Permanent and never reused, even if an account is deactivated.
        """
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        
        # Clean any users other than admin and John
        User.objects.exclude(id__in=[self.admin.id, self.employee.id]).delete()
        self.employee.employee_id = None
        self.employee.save()

        # Create Ayush
        res1 = self.client.post('/api/admin/users/', {
            'username': 'ayush_user1',
            'first_name': 'Ayush',
            'role': 'EMPLOYEE',
            'password': 'password123'
        })
        self.assertEqual(res1.status_code, 201)
        self.assertEqual(res1.data['user']['employee_id'], 'AYUSH001')

        # Create Aman
        res2 = self.client.post('/api/admin/users/', {
            'username': 'aman_user',
            'first_name': 'Aman',
            'role': 'EMPLOYEE',
            'password': 'password123'
        })
        self.assertEqual(res2.status_code, 201)
        self.assertEqual(res2.data['user']['employee_id'], 'AMAN002')

        # Create Ravina
        res3 = self.client.post('/api/admin/users/', {
            'username': 'ravina_user',
            'first_name': 'Ravina',
            'role': 'EMPLOYEE',
            'password': 'password123'
        })
        self.assertEqual(res3.status_code, 201)
        self.assertEqual(res3.data['user']['employee_id'], 'RAVINA003')

        # Create Amay
        res4 = self.client.post('/api/admin/users/', {
            'username': 'amay_user',
            'first_name': 'Amay',
            'role': 'EMPLOYEE',
            'password': 'password123'
        })
        self.assertEqual(res4.status_code, 201)
        self.assertEqual(res4.data['user']['employee_id'], 'AMAY004')

        # Create Saksham
        res5 = self.client.post('/api/admin/users/', {
            'username': 'saksham_user',
            'first_name': 'Saksham',
            'role': 'EMPLOYEE',
            'password': 'password123'
        })
        self.assertEqual(res5.status_code, 201)
        self.assertEqual(res5.data['user']['employee_id'], 'SAKSHAM005')

        # Create second Ayush (collision with same name) -> must use next available number (006)
        res6 = self.client.post('/api/admin/users/', {
            'username': 'ayush_user2',
            'first_name': 'Ayush',
            'role': 'EMPLOYEE',
            'password': 'password123'
        })
        self.assertEqual(res6.status_code, 201)
        self.assertEqual(res6.data['user']['employee_id'], 'AYUSH006')

        # Deactivate Ayush 1 (is_active=False)
        ayush1 = User.objects.get(username='ayush_user1')
        ayush1.is_active = False
        ayush1.save()

        # Create third Ayush -> must NOT reuse AYUSH001, must generate AYUSH007
        res7 = self.client.post('/api/admin/users/', {
            'username': 'ayush_user3',
            'first_name': 'Ayush',
            'role': 'EMPLOYEE',
            'password': 'password123'
        })
        self.assertEqual(res7.status_code, 201)
        self.assertEqual(res7.data['user']['employee_id'], 'AYUSH007')

    def test_admin_updates_employee_name_keeps_employee_id_unchanged(self):
        """
        If an employee's name changes, admin can update displayed name while keeping Employee ID unchanged.
        """
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        
        emp = User.objects.create_user(
            username='robert_u',
            first_name='Robert',
            last_name='Baratheon',
            role='EMPLOYEE',
            employee_id='ROBERT001'
        )

        res = self.client.patch(f'/api/admin/users/{emp.id}/', {
            'first_name': 'Robb',
            'last_name': 'Stark'
        })
        self.assertEqual(res.status_code, 200)

        emp.refresh_from_db()
        self.assertEqual(emp.first_name, 'Robb')
        self.assertEqual(emp.last_name, 'Stark')
        self.assertEqual(emp.employee_id, 'ROBERT001') # unchanged

    def test_task_creation_auto_attaches_employee_id_and_ignores_client_tampering(self):
        """
        When employee creates task, Employee ID and saved name are automatically attached
        from profile, and any client-provided name/id is ignored.
        """
        emp = User.objects.create_user(
            username='alice_w',
            first_name='Alice',
            last_name='Wonder',
            role='EMPLOYEE',
            employee_id='ALICE001',
            password='password123'
        )
        token, _ = Token.objects.get_or_create(user=emp)
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {token.key}')

        # Attempt to tamper by passing fake employee_id and employee_name in payload
        res = self.client.post('/api/tasks/start/', {
            'project_name': 'Project Phoenix',
            'task_name': 'Code Review',
            'description': 'Reviewing PR #42',
            'employee_id': 'HACKER999',
            'employee_name': 'Attacker'
        })
        self.assertEqual(res.status_code, 201)
        created_task = res.data['task']
        
        # Verify server used authenticated user's real employee_id and name
        self.assertEqual(created_task['employee_id'], 'ALICE001')
        self.assertEqual(created_task['employee_name'], 'Alice Wonder')

    def test_employee_can_only_view_own_tasks_admin_can_view_all(self):
        """
        Employees can view only their own task records. Admins can view all.
        """
        emp2 = User.objects.create_user(
            username='bob_builder',
            first_name='Bob',
            last_name='Builder',
            role='EMPLOYEE',
            employee_id='BOB001',
            password='password123'
        )
        token2, _ = Token.objects.get_or_create(user=emp2)

        # self.employee attempts to access emp2's tasks -> 403 Forbidden
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        res_emp = self.client.get(f'/api/tasks/daily/?user_id={emp2.id}')
        self.assertEqual(res_emp.status_code, 403)

        # Admin accesses emp2's tasks -> 200 OK
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        res_admin = self.client.get(f'/api/tasks/daily/?user_id={emp2.id}')
        self.assertEqual(res_admin.status_code, 200)
        self.assertEqual(res_admin.data['employee']['employee_id'], 'BOB001')

    def test_csv_and_excel_exports_contain_employee_id(self):
        """
        Verify Employee ID is included in admin reports export, tasks export, and timesheet export.
        """
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')

        # 1. Reports CSV Export
        res_rep_csv = self.client.get('/api/admin/reports/export/?format=csv')
        self.assertEqual(res_rep_csv.status_code, 200)
        self.assertIn('Employee ID', res_rep_csv.content.decode('utf-8'))

        # 2. Reports Excel Export
        res_rep_xlsx = self.client.get('/api/admin/reports/export/?format=xlsx')
        self.assertEqual(res_rep_xlsx.status_code, 200)
        self.assertEqual(res_rep_xlsx['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        # 3. Tasks CSV Export
        res_task_csv = self.client.get('/api/admin/tasks/export/?format=csv')
        self.assertEqual(res_task_csv.status_code, 200)
        self.assertIn('Employee ID', res_task_csv.content.decode('utf-8'))

        # 4. Tasks Excel Export
        res_task_xlsx = self.client.get('/api/admin/tasks/export/?format=xlsx')
        self.assertEqual(res_task_xlsx.status_code, 200)
        self.assertEqual(res_task_xlsx['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        # 5. Timesheet CSV Export
        res_ts_csv = self.client.get(f'/api/timesheet/export/?user_id={self.employee.id}&format=csv')
        self.assertEqual(res_ts_csv.status_code, 200)
        self.assertIn('Employee ID:', res_ts_csv.content.decode('utf-8'))

    def test_admin_cannot_track_employee_work(self):
        """
        Admins cannot record employee work sessions or log employee tasks.
        """
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        
        # Admin trying to start workday
        res_action = self.client.post('/api/tracker/action/', {'action': 'START_WORK'})
        self.assertEqual(res_action.status_code, 403)
        self.assertIn('Only employee accounts can track work time', res_action.data['error'])

        # Admin trying to start task
        res_task = self.client.post('/api/tasks/start/', {
            'project_name': 'Admin Work',
            'task_name': 'Reviewing logs',
            'description': 'Supervising'
        })
        self.assertEqual(res_task.status_code, 403)
        self.assertIn('Only employee accounts can log tasks', res_task.data['error'])

    def test_individual_employee_report_filter_and_export_verification(self):
        """
        Verify:
        - Admin can filter report by individual employee and date range / month.
        - The report rows and downloads contain only that employee's records.
        - Records include dates, start/end times, break duration, total work time, project, task, and description.
        - Export totals match the filtered report shown on screen.
        - Changing/clearing filter updates data.
        - Employees cannot download other employees' reports or access admin exports.
        """
        self.employee.employee_id = 'EMP001'
        self.employee.save()

        emp2 = User.objects.create_user(
            username='clara_oswald',
            first_name='Clara',
            last_name='Oswald',
            role='EMPLOYEE',
            employee_id='CLARA001',
            password='password123'
        )
        token2, _ = Token.objects.get_or_create(user=emp2)

        today_str = datetime.now(timezone.utc).astimezone(IST_TZ).date().isoformat()

        # Employee 1 records workday and a timesheet entry
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        self.client.post('/api/tracker/action/', {'action': 'START_WORK'})
        self.client.post('/api/timesheet/entry/', {
            'date': today_str,
            'project_name': 'Knowletive Cloud',
            'task_name': 'Core API Architecture',
            'work_description': 'Refining REST endpoints and authorization'
        })

        # Employee 2 records workday
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {token2.key}')
        self.client.post('/api/tracker/action/', {'action': 'START_WORK'})
        self.client.post('/api/timesheet/entry/', {
            'date': today_str,
            'project_name': 'Design System',
            'task_name': 'Component Styling',
            'work_description': 'Glassmorphism and light theme polish'
        })

        # 1. Employee cannot access Admin reports or export
        res_emp_rep = self.client.get('/api/admin/reports/')
        self.assertEqual(res_emp_rep.status_code, 403)
        res_emp_exp = self.client.get('/api/admin/reports/export/?format=csv')
        self.assertEqual(res_emp_exp.status_code, 403)

        # 2. Employee cannot tamper with user_id parameter on timesheet export
        # emp2 downloading own timesheet succeeds
        res_emp2_own = self.client.get(f'/api/timesheet/export/?user_id={emp2.id}&format=csv')
        self.assertEqual(res_emp2_own.status_code, 200)
        # But emp2 tampering with user_id to download Employee 1's timesheet gives 403 Forbidden
        res_tamper_other = self.client.get(f'/api/timesheet/export/?user_id={self.employee.id}&format=csv')
        self.assertEqual(res_tamper_other.status_code, 403)

        # 3. Admin filters specifically for Employee 1
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        res_admin_rep = self.client.get(f'/api/admin/reports/?employee_id={self.employee.id}&start_date={today_str}&end_date={today_str}')
        self.assertEqual(res_admin_rep.status_code, 200)
        rep_data = res_admin_rep.data['reports']
        summary = res_admin_rep.data['summary']

        # Contains ONLY Employee 1's records
        self.assertEqual(len(rep_data), 1)
        row = rep_data[0]
        self.assertEqual(row['employee_id'], self.employee.employee_id)
        self.assertEqual(row['project'], 'Knowletive Cloud')
        self.assertEqual(row['task'], 'Core API Architecture')
        self.assertEqual(row['description'], 'Refining REST endpoints and authorization')
        self.assertEqual(summary['total_records'], 1)

        # 4. Admin downloads CSV for Employee 1
        res_csv = self.client.get(f'/api/admin/reports/export/?format=csv&employee_id={self.employee.id}&start_date={today_str}&end_date={today_str}')
        self.assertEqual(res_csv.status_code, 200)
        csv_text = res_csv.content.decode('utf-8')
        self.assertIn(self.employee.employee_id, csv_text)
        self.assertNotIn('CLARA001', csv_text) # Other employee MUST NOT be present
        self.assertIn('Knowletive Cloud', csv_text)
        self.assertIn('Core API Architecture', csv_text)
        self.assertIn('TOTALS', csv_text)
        self.assertIn(f"{summary['total_work_hours']}", csv_text)

        # 5. Admin downloads Excel for Employee 1
        res_xlsx = self.client.get(f'/api/admin/reports/export/?format=xlsx&employee_id={self.employee.id}&start_date={today_str}&end_date={today_str}')
        self.assertEqual(res_xlsx.status_code, 200)
        stream = io.BytesIO(res_xlsx.content)
        wb = openpyxl.load_workbook(stream)
        ws = wb.active
        # Header row 4, Data row 5, Totals row 6
        cell_vals = [cell.value for cell in ws[5]]
        self.assertIn(self.employee.employee_id, cell_vals)
        self.assertIn('Knowletive Cloud', cell_vals)
        self.assertNotIn('CLARA001', [c.value for row in ws.rows for c in row if c.value])

        # 6. Clearing/changing filter to Employee 2
        res_admin_rep2 = self.client.get(f'/api/admin/reports/?employee_id={emp2.id}&start_date={today_str}&end_date={today_str}')
        self.assertEqual(res_admin_rep2.status_code, 200)
        self.assertEqual(res_admin_rep2.data['reports'][0]['employee_id'], 'CLARA001')
        self.assertEqual(res_admin_rep2.data['reports'][0]['project'], 'Design System')

    def test_public_signup_role_integrity(self):
        """
        Public employee signup must always create an Employee account.
        Users cannot choose or assign themselves Admin.
        """
        res = self.client.post('/api/auth/signup/', {
            'username': 'attacker_admin',
            'password': 'StrongPassword123!',
            'role': 'ADMIN' # Malicious attempt to elevate role
        })
        self.assertEqual(res.status_code, 201)
        created_user = User.objects.get(username='attacker_admin')
        self.assertEqual(created_user.role, 'EMPLOYEE')
        self.assertFalse(created_user.is_staff)
        self.assertFalse(created_user.is_superuser)
        self.assertTrue(created_user.employee_id.startswith('ATTACKER'))

    def test_admin_edit_user_work_email(self):
        """
        Verify that Admin can edit a user's work email in User Management,
        validations are enforced, and the user can subsequently sign in with the new email.
        """
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        
        # 1. Update employee's work email
        res = self.client.patch(f'/api/admin/users/{self.employee.id}/', {
            'email': 'new.work.email@knowletive.com'
        })
        self.assertEqual(res.status_code, 200)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.email, 'new.work.email@knowletive.com')

        # 2. Verify user list includes new email
        list_res = self.client.get('/api/admin/users/')
        self.assertEqual(list_res.status_code, 200)
        emp_data = next((u for u in list_res.data['users'] if u['id'] == self.employee.id), None)
        self.assertIsNotNone(emp_data)
        self.assertEqual(emp_data['email'], 'new.work.email@knowletive.com')

        # 3. Employee can log in with new work email
        self.client.credentials()  # Clear admin credentials for public login call
        login_res = self.client.post('/api/auth/login/', {
            'username': 'new.work.email@knowletive.com',
            'password': 'EmployeePassword123!'
        })
        self.assertEqual(login_res.status_code, 200)
        self.assertEqual(login_res.data['user']['email'], 'new.work.email@knowletive.com')

        # 4. Attempt to set duplicate email -> 400
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')
        dup_res = self.client.patch(f'/api/admin/users/{self.employee.id}/', {
            'email': self.admin.email
        })
        self.assertEqual(dup_res.status_code, 400)
        self.assertIn('already used', dup_res.data['error'])

        # 5. Attempt to set invalid email -> 400
        inv_res = self.client.patch(f'/api/admin/users/{self.employee.id}/', {
            'email': 'not-an-email'
        })
        self.assertEqual(inv_res.status_code, 400)
        self.assertIn('valid work email', inv_res.data['error'])

        # 6. Non-admin cannot edit user
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.employee_token.key}')
        unauth_res = self.client.patch(f'/api/admin/users/{self.admin.id}/', {
            'email': 'hacked@knowletive.com'
        })
        self.assertEqual(unauth_res.status_code, 403)


class AnalyticsDashboardTests(TestCase):
    def setUp(self):
        storage._ensure_mock_file()
        with open(storage.local_file, 'w', encoding='utf-8') as f:
            json.dump({
                'TimeEvents': [],
                'TimesheetEntries': [],
                'TaskLogs': [],
                'AuditLog': [],
                'Settings': {'work_start_time': '10:00', 'work_end_time': '19:00', 'timezone': 'Asia/Kolkata'},
                'Users': []
            }, f, indent=2)

        self.client = APIClient()
        self.admin = User.objects.create_user(
            username='admin_analytics@knowletive.com',
            email='admin_analytics@knowletive.com',
            password='Password123!',
            role='ADMIN',
            employee_id='ADM001',
            first_name='Admin',
            last_name='User'
        )
        self.admin_token = Token.objects.create(user=self.admin)

        self.emp1 = User.objects.create_user(
            username='emp1_analytics@knowletive.com',
            email='emp1_analytics@knowletive.com',
            password='Password123!',
            role='EMPLOYEE',
            employee_id='EMP001',
            first_name='Aman',
            last_name='Bendkule',
            department='SD'
        )
        self.emp1_token = Token.objects.create(user=self.emp1)

        self.emp2 = User.objects.create_user(
            username='emp2_analytics@knowletive.com',
            email='emp2_analytics@knowletive.com',
            password='Password123!',
            role='EMPLOYEE',
            employee_id='EMP002',
            first_name='Ayush',
            last_name='Kadam',
            department='AI/DATA'
        )
        self.emp2_token = Token.objects.create(user=self.emp2)

        # Seed realistic events
        storage.record_event(self.emp1.id, 'START_WORK')
        storage.record_event(self.emp1.id, 'START_BREAK')
        storage.record_event(self.emp1.id, 'RESUME_WORK')
        storage.record_event(self.emp1.id, 'END_WORK')

        # Seed tasks
        t1 = storage.start_task(self.emp1.id, 'Project Alpha', 'Feature 1', 'Desc', 'EMP001', 'Aman Bendkule')
        storage.end_task(self.emp1.id, t1['task_log_id'])

        t2 = storage.start_task(self.emp2.id, 'Project Beta', 'Model training', 'Desc', 'EMP002', 'Ayush Kadam')
        storage.end_task(self.emp2.id, t2['task_log_id'])

    def test_employee_permissions_enforced(self):
        # Authenticated as emp1
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.emp1_token.key}')
        
        # Request analytics
        res = self.client.get('/api/analytics/dashboard/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['user_role'], 'EMPLOYEE')
        self.assertFalse(res.data['is_admin'])

        # emp1 can only see their own tasks
        self.assertEqual(res.data['summary']['total_tasks'], 1)
        self.assertEqual(res.data['by_project'][0]['project_name'], 'Project Alpha')

        # If emp1 maliciously passes employee_id=EMP002, backend MUST enforce emp1's scope
        res_malicious = self.client.get(f'/api/analytics/dashboard/?employee_id={self.emp2.id}')
        self.assertEqual(res_malicious.status_code, 200)
        self.assertEqual(res_malicious.data['summary']['total_tasks'], 1)
        self.assertEqual(res_malicious.data['by_project'][0]['project_name'], 'Project Alpha')

    def test_admin_analytics_and_filtering(self):
        # Authenticated as admin
        self.client.credentials(HTTP_AUTHORIZATION=f'Token {self.admin_token.key}')

        # 1. Admin gets all team analytics
        res = self.client.get('/api/analytics/dashboard/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['user_role'], 'ADMIN')
        self.assertTrue(res.data['is_admin'])
        self.assertEqual(res.data['summary']['total_tasks'], 2)
        self.assertEqual(len(res.data['by_project']), 2)
        self.assertIn('Project Alpha', [p['project_name'] for p in res.data['by_project']])
        self.assertIn('Project Beta', [p['project_name'] for p in res.data['by_project']])

        # 2. Admin filters by employee
        res_filtered_emp = self.client.get(f'/api/analytics/dashboard/?employee_id={self.emp2.id}')
        self.assertEqual(res_filtered_emp.status_code, 200)
        self.assertEqual(res_filtered_emp.data['summary']['total_tasks'], 1)
        self.assertEqual(res_filtered_emp.data['by_project'][0]['project_name'], 'Project Beta')

        # 3. Admin filters by project
        res_filtered_proj = self.client.get('/api/analytics/dashboard/?project=Alpha')
        self.assertEqual(res_filtered_proj.status_code, 200)
        self.assertEqual(res_filtered_proj.data['summary']['total_tasks'], 1)
        self.assertEqual(res_filtered_proj.data['by_project'][0]['project_name'], 'Project Alpha')





