import csv
import io
import re
import secrets
import calendar
import zoneinfo
from datetime import datetime, timezone, timedelta, date
from django.db import models
from django.contrib.auth import authenticate
from django.http import HttpResponse
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.response import Response
from rest_framework.authtoken.models import Token
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from .models import User, Invitation, Notification, generate_unique_employee_id
from .storage_adapter import storage

IST_TZ = zoneinfo.ZoneInfo("Asia/Kolkata")


def is_valid_email(email_str):
    if not email_str or '@' not in email_str:
        return False
    # RFC 5322 simplified email regex pattern
    pattern = r'^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$'
    return bool(re.match(pattern, email_str))


def validate_strong_password(pw):
    if not pw or len(pw) < 8:
        return False, "Password must be at least 8 characters long."
    if not re.search(r'[A-Z]', pw):
        return False, "Password must contain at least one uppercase letter (A-Z)."
    if not re.search(r'[a-z]', pw):
        return False, "Password must contain at least one lowercase letter (a-z)."
    if not re.search(r'[0-9]', pw):
        return False, "Password must contain at least one number (0-9)."
    if not re.search(r'[!@#$%^&*(),.?":{}|<>_\-+=\[\]\\/`~]', pw):
        return False, "Password must contain at least one special character (!@#$%^&* etc.)."
    return True, ""


def parse_iso_utc(ts_str):
    if not ts_str:
        return None
    cleaned = ts_str.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(cleaned)
    except Exception:
        return None


def calculate_day_metrics(events_for_day, now_utc=None):
    if now_utc is None:
        now_utc = datetime.now(timezone.utc)
    
    events_sorted = sorted(events_for_day, key=lambda x: x['timestamp'])
    
    if not events_sorted:
        return {
            'status': 'Not Started',
            'tracker_state': 'IDLE',
            'start_time': None,
            'end_time': None,
            'total_work_seconds': 0,
            'total_break_seconds': 0,
            'completed_work_seconds': 0,
            'completed_break_seconds': 0,
            'work_duration': '00:00:00',
            'break_duration': '00:00:00',
            'total_daily_seconds': 0,
            'total_daily_duration': '00:00:00',
            'break_segments': [],
            'current_segment_start': None,
            'events': []
        }
    
    start_time = events_sorted[0]['timestamp']
    end_time = None
    total_work_seconds = 0
    total_break_seconds = 0
    
    current_state = 'Not Started'
    state_start_dt = None
    break_segments = []
    
    for ev in events_sorted:
        ev_dt = parse_iso_utc(ev['timestamp'])
        ev_type = ev['event_type']
        
        if current_state == 'Working':
            if state_start_dt and ev_dt:
                total_work_seconds += max(0, (ev_dt - state_start_dt).total_seconds())
        elif current_state == 'On Break':
            if state_start_dt and ev_dt:
                dur = max(0, (ev_dt - state_start_dt).total_seconds())
                total_break_seconds += dur
                break_segments.append({
                    'start_time': state_start_dt,
                    'end_time': ev_dt,
                    'duration_seconds': int(dur)
                })
                
        if ev_type in ['START_WORK', 'RESUME_WORK']:
            current_state = 'Working'
            state_start_dt = ev_dt
        elif ev_type == 'START_BREAK':
            current_state = 'On Break'
            state_start_dt = ev_dt
        elif ev_type == 'END_WORK':
            current_state = 'Ended'
            end_time = ev['timestamp']
            state_start_dt = None

    completed_work_sec = int(total_work_seconds)
    completed_break_sec = int(total_break_seconds)

    # If currently active, add duration up to now
    if current_state == 'Working' and state_start_dt:
        total_work_seconds += max(0, (now_utc - state_start_dt).total_seconds())
    elif current_state == 'On Break' and state_start_dt:
        dur = max(0, (now_utc - state_start_dt).total_seconds())
        total_break_seconds += dur
        break_segments.append({
            'start_time': state_start_dt,
            'end_time': None,
            'duration_seconds': int(dur)
        })

    # Format break details (start - end in IST)
    break_detail_items = []
    for b in break_segments:
        s_ist = b['start_time'].astimezone(IST_TZ).strftime('%I:%M:%S %p') if b['start_time'] else ''
        e_ist = b['end_time'].astimezone(IST_TZ).strftime('%I:%M:%S %p') if b['end_time'] else 'In progress'
        break_detail_items.append(f"{s_ist} - {e_ist}")
    break_details_str = "; ".join(break_detail_items) if break_detail_items else '-'

    # Format durations as HH:MM:SS
    total_b_sec = int(total_break_seconds)
    b_h = total_b_sec // 3600
    b_m = (total_b_sec % 3600) // 60
    b_s = total_b_sec % 60
    break_duration_formatted = f"{b_h:02d}:{b_m:02d}:{b_s:02d}"

    total_w_sec = int(total_work_seconds)
    w_h = total_w_sec // 3600
    w_m = (total_w_sec % 3600) // 60
    w_s = total_w_sec % 60
    work_duration_formatted = f"{w_h:02d}:{w_m:02d}:{w_s:02d}"

    # Calculate Total Daily Time: Clock-Out (or now_utc) - Clock-In
    total_daily_seconds = 0
    if start_time:
        start_dt_iso = parse_iso_utc(start_time)
        if start_dt_iso:
            if current_state in ['Working', 'On Break']:
                # currently in progress, use now_utc
                total_daily_seconds = max(0, (now_utc - start_dt_iso).total_seconds())
            elif end_time:
                end_dt_iso = parse_iso_utc(end_time)
                if end_dt_iso:
                    total_daily_seconds = max(0, (end_dt_iso - start_dt_iso).total_seconds())

    td_sec = int(total_daily_seconds)
    td_h = td_sec // 3600
    td_m = (td_sec % 3600) // 60
    td_s = td_sec % 60
    total_daily_duration_formatted = f"{td_h:02d}:{td_m:02d}:{td_s:02d}"

    tracker_state = 'WORKING' if current_state == 'Working' else ('BREAK' if current_state == 'On Break' else 'IDLE')

    return {
        'status': current_state,
        'tracker_state': tracker_state,
        'start_time': start_time,
        'end_time': end_time,
        'total_work_seconds': total_w_sec,
        'total_break_seconds': total_b_sec,
        'completed_work_seconds': completed_work_sec,
        'completed_break_seconds': completed_break_sec,
        'work_duration': work_duration_formatted,
        'break_duration': break_duration_formatted,
        'total_daily_seconds': td_sec,
        'total_daily_duration': total_daily_duration_formatted,
        'break_segments': [
            {
                'start_time': b['start_time'].isoformat() if b['start_time'] else None,
                'end_time': b['end_time'].isoformat() if b['end_time'] else None,
                'duration_seconds': b['duration_seconds']
            } for b in break_segments
        ],
        'current_segment_start': state_start_dt.isoformat().replace("+00:00", "Z") if state_start_dt else None,
        'events': events_sorted
    }


# ==========================================
# AUTHENTICATION
# ==========================================

@api_view(['POST'])
@permission_classes([AllowAny])
def login_view(request):
    username = (request.data.get('username') or request.data.get('email') or '').strip()
    password = request.data.get('password')
    
    if not username or not password:
        return Response({'error': 'Username and password are required.'}, status=400)
    
    # Authenticate by username (or fallback to email if someone entered email)
    user = User.objects.filter(username__iexact=username).first()
    if not user:
        user = User.objects.filter(email__iexact=username).first()
        
    if not user or not user.check_password(password):
        return Response({'error': 'Invalid username or password.'}, status=400)
        
    if not user.is_active:
        return Response({'error': 'Your account has been deactivated. Please contact an administrator.'}, status=403)
        
    token, _ = Token.objects.get_or_create(user=user)
    storage.log_audit(
        action="USER_LOGIN",
        actor=user.username,
        details={"user_id": user.id, "role": user.role}
    )
    return Response({
        'token': token.key,
        'user': {
            'id': user.id,
            'username': user.username,
            'employee_id': user.employee_id or '',
            'name': f"{user.first_name} {user.last_name}".strip() or user.username,
            'email': user.email or '',
            'role': user.role,
            'department': user.department or 'General',
            'employment_status': user.employment_status,
            'approval_status': user.approval_status
        }
    })


@api_view(['POST'])
@permission_classes([AllowAny])
def signup_view(request):
    username = (request.data.get('username') or '').strip()
    password = request.data.get('password') or ''

    if not username:
        return Response({'error': 'Username is required.'}, status=400)

    if len(username) < 3:
        return Response({'error': 'Username must be at least 3 characters long.'}, status=400)

    if not re.match(r'^[a-zA-Z0-9_.-]+$', username):
        return Response({'error': 'Username can only contain letters, numbers, dots, hyphens, and underscores.'}, status=400)

    if not password:
        return Response({'error': 'Password is required.'}, status=400)

    if len(password) < 6:
        return Response({'error': 'Password must be at least 6 characters long.'}, status=400)

    # Check for duplicate username
    if User.objects.filter(username__iexact=username).exists():
        return Response({'error': f'The username "{username}" is already taken. Please choose another username.'}, status=400)

    # Security requirement: ALL public signups are strictly assigned EMPLOYEE role.
    # Public visitors and signups CANNOT create or request an Admin account.
    emp_id = generate_unique_employee_id(username)
    user = User.objects.create_user(
        username=username,
        password=password,
        role='EMPLOYEE',
        employee_id=emp_id,
        approval_status='APPROVED',
        is_active=True,
        is_staff=False,
        is_superuser=False
    )

    token, _ = Token.objects.get_or_create(user=user)

    storage.log_audit(
        action="USER_SIGNUP",
        actor=user.username,
        details={
            "user_id": user.id,
            "username": user.username,
            "employee_id": user.employee_id,
            "role": "EMPLOYEE"
        }
    )

    return Response({
        'message': 'Account created successfully! You can now sign in.',
        'token': token.key,
        'user': {
            'id': user.id,
            'username': user.username,
            'employee_id': user.employee_id,
            'name': user.username,
            'role': 'EMPLOYEE',
            'approval_status': user.approval_status
        }
    }, status=201)


@api_view(['POST'])
@permission_classes([AllowAny])
def forgot_password_view(request):
    email = (request.data.get('email') or '').strip().lower()
    if not email:
        return Response({'error': 'Please enter your work email address.'}, status=400)
    if not is_valid_email(email):
        return Response({'error': 'Please enter a valid email address.'}, status=400)

    user = User.objects.filter(email__iexact=email).first()
    if user:
        Notification.objects.create(
            title="Password Reset Request",
            message=f"User {user.first_name} {user.last_name} ({user.email}) requested a password reset.",
            notification_type="PASSWORD_RESET",
            related_user=user
        )
        storage.log_audit(
            action="PASSWORD_RESET_REQUEST",
            actor=email,
            details={"user_id": user.id, "email": email}
        )

    return Response({
        'message': (
            'If your email is registered with Knowletive, an administrator notification has been recorded. '
            'Because direct email delivery is restricted for internal workspace security, '
            'please contact your workspace administrator to set or reset your credentials.'
        )
    })


@api_view(['GET'])
@permission_classes([AllowAny])
def verify_invitation_view(request):
    code = (request.GET.get('code') or '').strip().upper()
    if not code:
        return Response({'valid': False, 'error': 'Invitation code is required.'}, status=400)
    
    invitation = Invitation.objects.filter(code__iexact=code).first()
    if not invitation:
        return Response({'valid': False, 'error': 'Invalid invitation code.'}, status=404)
    if invitation.is_revoked:
        return Response({'valid': False, 'error': 'This invitation code has been revoked.'}, status=400)
    if invitation.is_used:
        return Response({'valid': False, 'error': 'This invitation code has already been used.'}, status=400)
    if invitation.expires_at and datetime.now(timezone.utc) > invitation.expires_at:
        return Response({'valid': False, 'error': 'This invitation code has expired.'}, status=400)

    return Response({
        'valid': True,
        'code': invitation.code,
        'recipient_email': invitation.recipient_email or '',
        'expires_at': invitation.expires_at.isoformat() if invitation.expires_at else None,
        'notes': invitation.notes
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def logout_view(request):
    try:
        request.user.auth_token.delete()
    except Exception:
        pass
    storage.log_audit(
        action="USER_LOGOUT",
        actor=request.user.email,
        details={"user_id": request.user.id}
    )
    return Response({'message': 'Logged out successfully.'})


# ==========================================
# EMPLOYEE TRACKER
# ==========================================

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def status_view(request):
    now_utc = datetime.now(timezone.utc)
    now_ist = now_utc.astimezone(IST_TZ)
    today_ist_date = now_ist.date()
    
    events = storage.get_events_for_user(request.user.id)
    
    # Filter today's events in IST
    today_events = []
    for ev in events:
        dt = parse_iso_utc(ev['timestamp'])
        if dt and dt.astimezone(IST_TZ).date() == today_ist_date:
            today_events.append(ev)
            
    today_metrics = calculate_day_metrics(today_events, now_utc)

    # Active or paused task details
    raw_active = storage.get_active_task_for_user(request.user.id)
    active_task = None
    if raw_active:
        dur_sec = storage.calculate_task_duration(raw_active, now_utc)
        active_task = {
            'task_log_id': raw_active['task_log_id'],
            'employee_id': raw_active.get('employee_id') or getattr(request.user, 'employee_id', '') or f"EMP{request.user.id:03d}",
            'employee_name': raw_active.get('employee_name') or f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username,
            'project_name': raw_active['project_name'],
            'task_name': raw_active['task_name'],
            'description': raw_active.get('description', ''),
            'start_time': raw_active['start_time'],
            'end_time': raw_active.get('end_time'),
            'status': raw_active.get('status', 'IN_PROGRESS'),
            'duration_seconds': dur_sec,
            'duration_seconds_base': int(raw_active.get('duration_seconds') or 0),
            'current_segment_start': raw_active.get('current_segment_start'),
            'work_segments': raw_active.get('work_segments', [])
        }
    
    # Weekly history (past 7 days in IST)
    weekly_history = []
    for i in range(7):
        target_date = today_ist_date - timedelta(days=i)
        day_events = [
            e for e in events 
            if parse_iso_utc(e['timestamp']) and parse_iso_utc(e['timestamp']).astimezone(IST_TZ).date() == target_date
        ]
        metrics = calculate_day_metrics(day_events, now_utc if target_date == today_ist_date else None)
        weekly_history.append({
            'date': target_date.isoformat(),
            'status': metrics['status'],
            'start_time': metrics['start_time'],
            'end_time': metrics['end_time'],
            'total_work_seconds': metrics['total_work_seconds'],
            'total_break_seconds': metrics['total_break_seconds']
        })

    settings = storage.get_settings()
    
    return Response({
        'status': today_metrics['status'],
        'tracker_state': today_metrics['tracker_state'],
        'start_time': today_metrics['start_time'],
        'end_time': today_metrics['end_time'],
        'total_work_seconds': today_metrics['total_work_seconds'],
        'total_break_seconds': today_metrics['total_break_seconds'],
        'completed_work_seconds': today_metrics['completed_work_seconds'],
        'completed_break_seconds': today_metrics['completed_break_seconds'],
        'current_segment_start': today_metrics['current_segment_start'],
        'active_task': active_task,
        'today_events': today_metrics['events'],
        'weekly_history': weekly_history,
        'settings': settings
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def action_view(request):
    if request.user.role != 'EMPLOYEE':
        return Response({'error': 'Only employee accounts can track work time. Admin accounts cannot record employee work logs.'}, status=403)

    action = request.data.get('action')
    source = request.data.get('source', 'web')
    valid_actions = ['START_WORK', 'START_BREAK', 'RESUME_WORK', 'END_WORK']
    
    if action not in valid_actions:
        return Response({'error': f'Invalid action: {action}.'}, status=400)

    now_utc = datetime.now(timezone.utc)
    now_utc_str = now_utc.isoformat().replace("+00:00", "Z")
    now_ist = now_utc.astimezone(IST_TZ)
    today_ist_date = now_ist.date()
    
    # Fetch events & lock checks
    events = storage.get_events_for_user(request.user.id)
    today_events = [
        e for e in events 
        if parse_iso_utc(e['timestamp']) and parse_iso_utc(e['timestamp']).astimezone(IST_TZ).date() == today_ist_date
    ]
    today_events.sort(key=lambda x: x['timestamp'])

    # Determine current status
    metrics = calculate_day_metrics(today_events, now_utc)
    current_status = metrics['status']

    # State machine transition validation
    if action == 'START_WORK':
        if current_status == 'Working':
            return Response({'error': 'Workday is already active.'}, status=400)
        if current_status == 'On Break':
            return Response({'error': 'You are currently on a break. Please use Resume Work.'}, status=400)
        if current_status == 'Ended':
            return Response({'error': 'Workday has already ended for today. Contact admin for corrections.'}, status=400)
            
    elif action == 'START_BREAK':
        if current_status != 'Working':
            return Response({'error': 'You must be actively working to start a break.'}, status=400)
            
    elif action == 'RESUME_WORK':
        if current_status != 'On Break':
            return Response({'error': 'You are not currently on a break.'}, status=400)
            
    elif action == 'END_WORK':
        if current_status not in ['Working', 'On Break']:
            return Response({'error': 'Cannot end workday because it is not active.'}, status=400)

    # Deduplication check: if action is identical to last event within 2 seconds, ignore to protect against rapid double clicks
    if today_events:
        last_ev = today_events[-1]
        last_dt = parse_iso_utc(last_ev['timestamp'])
        if last_ev['event_type'] == action and (now_utc - last_dt).total_seconds() < 2:
            return Response({'error': 'Duplicate request ignored.'}, status=429)

    # CASE 2: When starting break, immediately PAUSE active task and freeze its work duration
    if action == 'START_BREAK':
        storage.pause_active_task(request.user.id, now_utc_str)

    # CASE 3: When resuming work from break, immediately RESUME the paused task
    elif action == 'RESUME_WORK':
        storage.resume_active_task(request.user.id, now_utc_str)

    # When ending workday, close any active or paused task
    elif action == 'END_WORK':
        storage.end_any_active_task_for_user(request.user.id, now_utc_str)

    # Server writes timestamp, ignores any client timestamp
    res = storage.record_event(request.user.id, action, source=source)
    return Response({'success': True, 'event': res})


# ==========================================
# ADMIN DASHBOARD & ATTENDANCE
# ==========================================

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_dashboard(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    users = User.objects.filter(role='EMPLOYEE').order_by('first_name', 'last_name')
    all_events = storage.get_all_events()
    settings = storage.get_settings()
    
    now_utc = datetime.now(timezone.utc)
    now_ist = now_utc.astimezone(IST_TZ)
    today_ist_date = now_ist.date()

    # Parse office hours for late/early comparison
    work_start_h, work_start_m = map(int, settings.get('work_start_time', '10:00').split(':'))
    work_end_h, work_end_m = map(int, settings.get('work_end_time', '19:00').split(':'))
    
    office_start_dt = now_ist.replace(hour=work_start_h, minute=work_start_m, second=0, microsecond=0)
    office_end_dt = now_ist.replace(hour=work_end_h, minute=work_end_m, second=0, microsecond=0)

    user_data = []
    active_count = 0
    late_count = 0
    missing_clockout_count = 0

    for u in users:
        u_events = [
            e for e in all_events 
            if int(e.get('user_id', -1)) == u.id and parse_iso_utc(e['timestamp']) and parse_iso_utc(e['timestamp']).astimezone(IST_TZ).date() == today_ist_date
        ]
        metrics = calculate_day_metrics(u_events, now_utc)
        
        status = metrics['status']
        if status in ['Working', 'On Break']:
            active_count += 1

        is_late_start = False
        if metrics['start_time']:
            start_dt = parse_iso_utc(metrics['start_time']).astimezone(IST_TZ)
            if start_dt > office_start_dt + timedelta(minutes=5):
                is_late_start = True
                late_count += 1

        is_early_finish = False
        if metrics['end_time']:
            end_dt = parse_iso_utc(metrics['end_time']).astimezone(IST_TZ)
            if end_dt < office_end_dt - timedelta(minutes=10):
                is_early_finish = True

        is_missing_clockout = False
        if status in ['Working', 'On Break'] and now_ist > office_end_dt + timedelta(hours=1):
            is_missing_clockout = True
            missing_clockout_count += 1

        user_data.append({
            'id': u.id,
            'employee_id': u.employee_id or f"EMP{u.id:03d}",
            'name': f"{u.first_name} {u.last_name}".strip() or u.username,
            'email': u.email,
            'department': u.department or 'Unassigned',
            'role': u.role,
            'employment_status': u.employment_status,
            'is_active': u.is_active,
            'status': status,
            'start_time': metrics['start_time'],
            'end_time': metrics['end_time'],
            'total_work_seconds': metrics['total_work_seconds'],
            'total_break_seconds': metrics['total_break_seconds'],
            'is_late_start': is_late_start,
            'is_early_finish': is_early_finish,
            'is_missing_clockout': is_missing_clockout
        })
        
    return Response({
        'overview': {
            'total_employees': users.count(),
            'active_now': active_count,
            'late_today': late_count,
            'missing_clockouts': missing_clockout_count
        },
        'users': user_data,
        'settings': settings
    })


# ==========================================
# USER MANAGEMENT (CRUD)
# ==========================================

@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def user_management(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    if request.method == 'GET':
        users = User.objects.all().order_by('id')
        data = [{
            'id': u.id,
            'employee_id': u.employee_id or (f"EMP{u.id:03d}" if u.role == 'EMPLOYEE' else ''),
            'name': f"{u.first_name} {u.last_name}".strip() or u.username,
            'first_name': u.first_name,
            'last_name': u.last_name,
            'email': u.email,
            'department': u.department,
            'role': u.role,
            'employment_status': u.employment_status,
            'approval_status': u.approval_status,
            'is_active': u.is_active,
            'date_joined': u.date_joined.isoformat() if u.date_joined else None
        } for u in users]
        return Response({'users': data})

    elif request.method == 'POST':
        username = (request.data.get('username') or request.data.get('email') or '').strip()
        email = (request.data.get('email') or '').strip().lower()
        full_name_input = (request.data.get('name') or '').strip()
        first_name = (request.data.get('first_name') or '').strip()
        last_name = (request.data.get('last_name') or '').strip()
        if full_name_input and not first_name:
            parts = full_name_input.split(' ', 1)
            first_name = parts[0]
            if len(parts) > 1 and not last_name:
                last_name = parts[1]

        department = (request.data.get('department') or '').strip()
        role = request.data.get('role', 'EMPLOYEE')
        password = request.data.get('password')
        approval_status = request.data.get('approval_status', 'APPROVED')
        
        if not username:
            return Response({'error': 'Username is required.'}, status=400)
            
        if User.objects.filter(username__iexact=username).exists():
            return Response({'error': f'A user with username "{username}" already exists.'}, status=400)

        if email:
            if not is_valid_email(email):
                return Response({'error': 'Please provide a valid work email address.'}, status=400)
            if User.objects.filter(email__iexact=email).exists():
                return Response({'error': f'The email "{email}" is already used by another account.'}, status=400)
            
        if password:
            if len(password) < 6:
                return Response({'error': 'Password must be at least 6 characters long.'}, status=400)
        else:
            password = f"K{secrets.token_urlsafe(8)}!9a"

        assigned_role = role if role in ['ADMIN', 'EMPLOYEE'] else 'EMPLOYEE'
        emp_id = None
        if assigned_role == 'EMPLOYEE':
            emp_id = generate_unique_employee_id(first_name or full_name_input or username)

        user = User.objects.create_user(
            username=username,
            email=email if email else f"{username}@knowletive.local",
            first_name=first_name,
            last_name=last_name,
            department=department,
            role=assigned_role,
            employee_id=emp_id,
            approval_status='APPROVED',
            is_active=True,
            is_staff=(assigned_role == 'ADMIN'),
            is_superuser=(assigned_role == 'ADMIN'),
            password=password
        )
        storage.log_audit(
            action="CREATE_USER",
            actor=request.user.username,
            details={"created_user_id": user.id, "username": username, "employee_id": user.employee_id, "role": user.role}
        )
        return Response({
            'message': 'User created successfully.',
            'user': {
                'id': user.id,
                'employee_id': user.employee_id or '',
                'username': user.username,
                'email': user.email,
                'name': f"{user.first_name} {user.last_name}".strip() or user.username,
                'role': user.role,
                'department': user.department,
                'approval_status': user.approval_status
            }
        }, status=201)


@api_view(['PATCH', 'PUT'])
@permission_classes([IsAuthenticated])
def user_detail_manage(request, user_id):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    try:
        user = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'User not found.'}, status=404)
        
    data = request.data
    previous_role = user.role

    if 'first_name' in data:
        user.first_name = data['first_name']
    if 'last_name' in data:
        user.last_name = data['last_name']
    if 'email' in data:
        new_email = (data.get('email') or '').strip().lower()
        if not new_email:
            return Response({'error': 'Work email cannot be blank.'}, status=400)
        if not is_valid_email(new_email):
            return Response({'error': 'Please provide a valid work email address.'}, status=400)
        if User.objects.filter(email__iexact=new_email).exclude(id=user.id).exists():
            return Response({'error': f'The email "{new_email}" is already used by another account.'}, status=400)
        
        old_email = (user.email or '').strip().lower()
        user.email = new_email
        # If user was using old email as username, sync username if available
        if old_email and user.username.lower() == old_email:
            if not User.objects.filter(username__iexact=new_email).exclude(id=user.id).exists():
                user.username = new_email

    if 'username' in data:
        new_username = (data.get('username') or '').strip()
        if new_username and new_username.lower() != user.username.lower():
            if User.objects.filter(username__iexact=new_username).exclude(id=user.id).exists():
                return Response({'error': f'The username "{new_username}" is already taken.'}, status=400)
            user.username = new_username

    if 'department' in data:
        user.department = data['department']
    if 'role' in data and data['role'] in ['ADMIN', 'EMPLOYEE']:
        new_role = data['role']
        if previous_role != new_role:
            storage.log_audit(
                action="ADMIN_ROLE_CHANGE",
                actor=request.user.email or request.user.username,
                details={
                    "target_user_id": user.id,
                    "target_email": user.email,
                    "previous_role": previous_role,
                    "new_role": new_role,
                    "changed_by": request.user.email or request.user.username
                }
            )
        user.role = new_role
        user.is_staff = (new_role == 'ADMIN')
        user.is_superuser = (new_role == 'ADMIN')
    if 'employment_status' in data:
        user.employment_status = data['employment_status']
    if 'approval_status' in data and data['approval_status'] in ['PENDING', 'APPROVED', 'REJECTED']:
        user.approval_status = data['approval_status']
    if 'is_active' in data:
        val = data['is_active']
        is_active_bool = val if isinstance(val, bool) else str(val).lower() in ['true', '1', 't']
        if not is_active_bool and user.role == 'ADMIN':
            return Response({'error': 'Administrator accounts cannot be deactivated.'}, status=400)
        user.is_active = is_active_bool
    if 'password' in data and data['password']:
        is_valid, pw_err = validate_strong_password(data['password'])
        if not is_valid:
            return Response({'error': f'Password does not meet complexity requirements: {pw_err}'}, status=400)
        user.set_password(data['password'])
        storage.log_audit(
            action="ADMIN_PASSWORD_RESET",
            actor=request.user.email or request.user.username,
            details={"target_user_id": user.id, "target_email": user.email}
        )
        
    # Keep employee_id permanent and generated if employee role has no id
    if user.role == 'EMPLOYEE' and not user.employee_id:
        user.employee_id = generate_unique_employee_id(user.first_name or user.username)
        
    user.save()
    storage.log_audit(
        action="UPDATE_USER",
        actor=request.user.email or request.user.username,
        details={
            "updated_user_id": user.id,
            "employee_id": user.employee_id,
            "email": user.email,
            "role": user.role,
            "approval_status": user.approval_status,
            "is_active": user.is_active
        }
    )
    return Response({
        'message': 'User updated successfully.',
        'user': {
            'id': user.id,
            'employee_id': user.employee_id or '',
            'name': f"{user.first_name} {user.last_name}".strip() or user.username,
            'email': user.email,
            'role': user.role
        }
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def approve_user_view(request, user_id):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
    try:
        user = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'User not found.'}, status=404)

    user.approval_status = 'APPROVED'
    user.is_active = True
    user.save()
    storage.log_audit(
        action="APPROVE_USER",
        actor=request.user.email,
        details={"approved_user_id": user.id, "email": user.email}
    )
    return Response({'message': f'User {user.email} has been approved successfully.'})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def reject_user_view(request, user_id):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
    try:
        user = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'User not found.'}, status=404)

    user.approval_status = 'REJECTED'
    user.save()
    storage.log_audit(
        action="REJECT_USER",
        actor=request.user.email,
        details={"rejected_user_id": user.id, "email": user.email}
    )
    return Response({'message': f'User {user.email} has been rejected.'})


# ==========================================
# INVITATION MANAGEMENT
# ==========================================

@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def admin_invitations_view(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    if request.method == 'GET':
        invitations = Invitation.objects.all().order_by('-created_at')
        now_dt = datetime.now(timezone.utc)
        data = []
        for inv in invitations:
            status_str = 'Active'
            if inv.is_revoked:
                status_str = 'Revoked'
            elif inv.is_used:
                status_str = 'Used'
            elif inv.expires_at and now_dt > inv.expires_at:
                status_str = 'Expired'

            data.append({
                'id': inv.id,
                'code': inv.code,
                'recipient_email': inv.recipient_email or '',
                'created_at': inv.created_at.isoformat() if inv.created_at else None,
                'expires_at': inv.expires_at.isoformat() if inv.expires_at else None,
                'is_used': inv.is_used,
                'used_by': f"{inv.used_by.first_name} {inv.used_by.last_name} ({inv.used_by.email})".strip() if inv.used_by else None,
                'used_at': inv.used_at.isoformat() if inv.used_at else None,
                'is_revoked': inv.is_revoked,
                'status': status_str,
                'notes': inv.notes or ''
            })
        return Response({'invitations': data})

    elif request.method == 'POST':
        recipient_email = (request.data.get('recipient_email') or '').strip().lower()
        days_valid = int(request.data.get('days_valid') or 7)
        notes = (request.data.get('notes') or '').strip()
        custom_code = (request.data.get('code') or '').strip().upper()

        if recipient_email and not is_valid_email(recipient_email):
            return Response({'error': 'Invalid recipient email format.'}, status=400)

        if custom_code:
            if Invitation.objects.filter(code__iexact=custom_code).exists():
                return Response({'error': 'An invitation with this code already exists.'}, status=400)
            code = custom_code
        else:
            code = f"KNOW-{secrets.token_hex(4).upper()}"
            while Invitation.objects.filter(code=code).exists():
                code = f"KNOW-{secrets.token_hex(4).upper()}"

        expires_at = datetime.now(timezone.utc) + timedelta(days=days_valid) if days_valid > 0 else None

        inv = Invitation.objects.create(
            code=code,
            created_by=request.user,
            recipient_email=recipient_email if recipient_email else None,
            expires_at=expires_at,
            notes=notes
        )

        storage.log_audit(
            action="CREATE_INVITATION",
            actor=request.user.email,
            details={
                "code": code,
                "recipient_email": recipient_email,
                "expires_at": expires_at.isoformat() if expires_at else None,
                "notes": notes
            }
        )

        return Response({
            'message': 'Invitation created successfully.',
            'invitation': {
                'id': inv.id,
                'code': inv.code,
                'recipient_email': inv.recipient_email or '',
                'expires_at': inv.expires_at.isoformat() if inv.expires_at else None,
                'status': 'Active',
                'notes': inv.notes
            }
        }, status=201)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def revoke_invitation_view(request, inv_id):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
    try:
        inv = Invitation.objects.get(id=inv_id)
    except Invitation.DoesNotExist:
        return Response({'error': 'Invitation not found.'}, status=404)

    inv.is_revoked = True
    inv.save()
    storage.log_audit(
        action="REVOKE_INVITATION",
        actor=request.user.email,
        details={"invitation_id": inv.id, "code": inv.code}
    )
    return Response({'message': f'Invitation {inv.code} revoked successfully.'})


# ==========================================
# ADMIN NOTIFICATIONS
# ==========================================

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_notifications_view(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
    notifications = Notification.objects.all().order_by('-created_at')[:40]
    data = [{
        'id': n.id,
        'title': n.title,
        'message': n.message,
        'notification_type': n.notification_type,
        'related_user_id': n.related_user_id,
        'created_at': n.created_at.isoformat(),
        'is_read': n.is_read
    } for n in notifications]
    unread_count = Notification.objects.filter(is_read=False).count()
    return Response({'notifications': data, 'unread_count': unread_count})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def mark_notifications_read_view(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
    notif_id = request.data.get('notification_id')
    if notif_id:
        Notification.objects.filter(id=notif_id).update(is_read=True)
    else:
        Notification.objects.filter(is_read=False).update(is_read=True)
    return Response({'message': 'Notifications marked as read.'})


# ==========================================
# ADMIN CORRECTIONS & AUDIT LOGS
# ==========================================

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def admin_correction(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    event_id = request.data.get('event_id')
    new_event_type = request.data.get('new_event_type')
    new_timestamp = request.data.get('new_timestamp')
    reason = (request.data.get('reason') or '').strip()
    
    if not event_id:
        return Response({'error': 'Event ID is required.'}, status=400)
    if not reason:
        return Response({'error': 'A detailed reason for correction is strictly required for audit compliance.'}, status=400)

    try:
        storage.correct_event(
            event_id=event_id,
            new_event_type=new_event_type,
            new_timestamp=new_timestamp,
            admin_email=request.user.email,
            reason=reason
        )
        return Response({'message': 'Correction applied and audit logged successfully.'})
    except ValueError as e:
        return Response({'error': str(e)}, status=400)
    except Exception as e:
        return Response({'error': f'Failed to apply correction: {str(e)}'}, status=500)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def audit_logs_view(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    logs = storage.get_audit_logs()
    return Response({'audit_logs': logs})


# ==========================================
# REPORTS & EXPORTS (CSV / XLSX)
# ==========================================

def get_filtered_report_data(request):
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    month = request.GET.get('month')
    year = request.GET.get('year')
    employee_id = request.GET.get('employee_id') or request.GET.get('user_id')
    department = request.GET.get('department')
    filter_project = request.GET.get('project')
    filter_task = request.GET.get('task')
    
    users = User.objects.filter(role='EMPLOYEE')
    if employee_id:
        if str(employee_id).isdigit():
            users = users.filter(models.Q(id=int(employee_id)) | models.Q(employee_id=employee_id))
        else:
            users = users.filter(employee_id=employee_id)
    if department:
        users = users.filter(department__iexact=department)
        
    all_events = storage.get_all_events()
    report_rows = []
    
    now_utc = datetime.now(timezone.utc)
    now_ist = now_utc.astimezone(IST_TZ)
    today_ist_date = now_ist.date()
    
    # Gather dates: support both explicit date range and month selection
    if month:
        try:
            m_val = int(month)
            y_val = int(year) if year else today_ist_date.year
            num_days = calendar.monthrange(y_val, m_val)[1]
            d_start = date(y_val, m_val, 1)
            d_end = date(y_val, m_val, num_days)
        except (ValueError, TypeError):
            d_start = today_ist_date - timedelta(days=6)
            d_end = today_ist_date
    elif start_date and end_date:
        try:
            d_start = datetime.strptime(start_date, '%Y-%m-%d').date()
            d_end = datetime.strptime(end_date, '%Y-%m-%d').date()
        except ValueError:
            d_start = today_ist_date - timedelta(days=6)
            d_end = today_ist_date
    elif start_date:
        try:
            d_start = datetime.strptime(start_date, '%Y-%m-%d').date()
            d_end = d_start
        except ValueError:
            d_start = today_ist_date - timedelta(days=6)
            d_end = today_ist_date
    else:
        d_start = today_ist_date - timedelta(days=6)
        d_end = today_ist_date

    # Pre-cache timesheet entries for users in date range
    entries_map = {}
    for u in users:
        cur_m_date = d_start.replace(day=1)
        end_m_date = d_end.replace(day=1)
        while cur_m_date <= end_m_date:
            u_entries = storage.get_timesheet_entries_for_month(u.id, cur_m_date.year, cur_m_date.month)
            for entry in u_entries:
                entries_map[(u.id, entry.get('date'))] = entry
            days_in_m = calendar.monthrange(cur_m_date.year, cur_m_date.month)[1]
            cur_m_date = cur_m_date + timedelta(days=days_in_m)

    # Iterate each day and each user
    curr = d_start
    while curr <= d_end:
        curr_str = curr.isoformat()
        for u in users:
            u_events = [
                e for e in all_events 
                if int(e.get('user_id', -1)) == u.id and parse_iso_utc(e['timestamp']) and parse_iso_utc(e['timestamp']).astimezone(IST_TZ).date() == curr
            ]
            if not u_events and (curr == today_ist_date or u.is_active):
                pass # Still may record Not Started
            metrics = calculate_day_metrics(u_events, now_utc if curr == today_ist_date else None)
            
            # Format times to IST
            start_str = '-'
            if metrics['start_time']:
                start_str = parse_iso_utc(metrics['start_time']).astimezone(IST_TZ).strftime('%I:%M:%S %p')
                
            end_str = '-'
            if metrics['end_time']:
                end_str = parse_iso_utc(metrics['end_time']).astimezone(IST_TZ).strftime('%I:%M:%S %p')
                
            work_hours = round(metrics['total_work_seconds'] / 3600.0, 2)
            break_hours = round(metrics['total_break_seconds'] / 3600.0, 2)

            # Retrieve project, task, description
            saved_entry = entries_map.get((u.id, curr_str), {})
            proj_val = saved_entry.get('project_name', '')
            task_val = saved_entry.get('task_name', '')
            desc_val = saved_entry.get('work_description', '')

            if not proj_val and not task_val:
                day_tasks = storage.get_tasks_for_user_and_date(u.id, curr_str)
                if day_tasks:
                    proj_val = ", ".join(dict.fromkeys(t['project_name'] for t in day_tasks if t.get('project_name')))
                    task_val = ", ".join(dict.fromkeys(t['task_name'] for t in day_tasks if t.get('task_name')))
                    desc_val = "; ".join(t['description'] for t in day_tasks if t.get('description'))

            if filter_project and filter_project.lower() not in proj_val.lower():
                continue
            if filter_task and filter_task.lower() not in task_val.lower():
                continue

            report_rows.append({
                'date': curr_str,
                'user_id': u.id,
                'employee_id': u.employee_id or f"EMP{u.id:03d}",
                'employee_name': f"{u.first_name} {u.last_name}".strip() or u.username,
                'email': u.email or '',
                'department': u.department or 'Unassigned',
                'project': proj_val or '-',
                'task': task_val or '-',
                'description': desc_val or '-',
                'status': metrics['status'],
                'start_time': start_str,
                'end_time': end_str,
                'total_work_seconds': metrics['total_work_seconds'],
                'total_break_seconds': metrics['total_break_seconds'],
                'work_hours': work_hours,
                'work_duration': metrics['work_duration'],
                'break_hours': break_hours,
                'break_duration': metrics['break_duration'],
                'total_daily_seconds': metrics.get('total_daily_seconds', 0),
                'total_daily_duration': metrics.get('total_daily_duration', '00:00:00')
            })
        curr += timedelta(days=1)
        
    return report_rows, d_start, d_end, users


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def reports_view(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    data, d_start, d_end, users = get_filtered_report_data(request)
    total_work_hours = round(sum(r['work_hours'] for r in data), 2)
    total_break_hours = round(sum(r['break_hours'] for r in data), 2)
    total_work_seconds = sum(r['total_work_seconds'] for r in data)
    total_break_seconds = sum(r['total_break_seconds'] for r in data)

    tw_h = total_work_seconds // 3600
    tw_m = (total_work_seconds % 3600) // 60
    tw_s = total_work_seconds % 60
    total_work_dur = f"{tw_h:02d}:{tw_m:02d}:{tw_s:02d}"

    tb_h = total_break_seconds // 3600
    tb_m = (total_break_seconds % 3600) // 60
    tb_s = total_break_seconds % 60
    total_break_dur = f"{tb_h:02d}:{tb_m:02d}:{tb_s:02d}"

    total_daily_seconds = sum(r.get('total_daily_seconds', 0) for r in data)
    td_h = total_daily_seconds // 3600
    td_m = (total_daily_seconds % 3600) // 60
    td_s = total_daily_seconds % 60
    total_daily_dur = f"{td_h:02d}:{td_m:02d}:{td_s:02d}"

    return Response({
        'reports': data,
        'summary': {
            'total_records': len(data),
            'total_work_hours': total_work_hours,
            'total_break_hours': total_break_hours,
            'total_work_seconds': total_work_seconds,
            'total_break_seconds': total_break_seconds,
            'total_daily_seconds': total_daily_seconds,
            'total_work_duration': total_work_dur,
            'total_break_duration': total_break_dur,
            'total_daily_duration': total_daily_dur
        }
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_reports_view(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    export_format = request.GET.get('format', 'csv').lower()
    data, d_start, d_end, users = get_filtered_report_data(request)
    total_work_hours = round(sum(r['work_hours'] for r in data), 2)
    total_break_hours = round(sum(r['break_hours'] for r in data), 2)
    total_work_seconds = sum(r['total_work_seconds'] for r in data)
    total_break_seconds = sum(r['total_break_seconds'] for r in data)

    tw_h = total_work_seconds // 3600
    tw_m = (total_work_seconds % 3600) // 60
    tw_s = total_work_seconds % 60
    total_work_dur = f"{tw_h:02d}:{tw_m:02d}:{tw_s:02d}"

    tb_h = total_break_seconds // 3600
    tb_m = (total_break_seconds % 3600) // 60
    tb_s = total_break_seconds % 60
    total_break_dur = f"{tb_h:02d}:{tb_m:02d}:{tb_s:02d}"
    
    total_daily_seconds_sum = sum(r.get('total_daily_seconds', 0) for r in data)
    td_h = total_daily_seconds_sum // 3600
    td_m = (total_daily_seconds_sum % 3600) // 60
    td_s = total_daily_seconds_sum % 60
    total_daily_dur = f"{td_h:02d}:{td_m:02d}:{td_s:02d}"
    
    # Audit export
    storage.log_audit(
        action="EXPORT_REPORT",
        actor=request.user.email,
        details={"format": export_format, "rows_count": len(data), "total_work_hours": total_work_hours}
    )

    # Determine filename
    start_str = d_start.strftime('%Y-%m-%d')
    end_str = d_end.strftime('%Y-%m-%d')
    date_range_str = f"{start_str}_to_{end_str}"
    
    employee_id_param = request.GET.get('employee_id')
    if employee_id_param and users.exists():
        emp_obj = users.first()
        raw_name = f"{emp_obj.first_name} {emp_obj.last_name}".strip() or emp_obj.username
        emp_name = re.sub(r'[^A-Za-z0-9]+', '_', raw_name).strip('_') or 'Employee'
    elif users.count() == 1:
        emp_obj = users.first()
        raw_name = f"{emp_obj.first_name} {emp_obj.last_name}".strip() or emp_obj.username
        emp_name = re.sub(r'[^A-Za-z0-9]+', '_', raw_name).strip('_') or 'Employee'
    else:
        emp_name = "All_Employees"

    filename = f"{emp_name}_Time_Report_{date_range_str}"

    headers = [
        'Date (IST)', 'Employee ID', 'Employee Name', 'Email', 
        'Department', 'Project', 'Task Name', 'Work Description', 'Status', 
        'Clock-In (IST)', 'Clock-Out (IST)', 
        'Work Duration', 'Break Duration', 'Total Daily Time'
    ]

    if export_format == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
        response['Access-Control-Expose-Headers'] = 'Content-Disposition, X-Suggested-Filename'
        response['X-Suggested-Filename'] = f"{filename}.csv"
        
        writer = csv.writer(response)
        writer.writerow(['KNOWLETIVE EMPLOYEE ATTENDANCE & TIME REPORT'])
        writer.writerow([
            'Total Records:', len(data), 
            'Total Work Duration:', total_work_dur, 
            'Total Break Duration:', total_break_dur
        ])
        writer.writerow([])
        writer.writerow(headers)
        for r in data:
            writer.writerow([
                r['date'], r['employee_id'], r['employee_name'], r['email'],
                r['department'], r['project'], r['task'], r['description'], r['status'], 
                r['start_time'], r['end_time'],
                r['work_duration'], r['break_duration'], r['total_daily_duration']
            ])
        writer.writerow([])
        writer.writerow([
            'TOTALS', '', '', '', '', '', '', '', '', '', '',
            total_work_dur, 
            total_break_dur, 
            total_daily_dur
        ])
        return response

    elif export_format in ['xlsx', 'excel']:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Time Report"
        
        title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        ws.append(['KNOWLETIVE EMPLOYEE ATTENDANCE & TIME REPORT'])
        ws['A1'].font = title_font
        ws.append([
            f"Total Records: {len(data)}", 
            f"Total Work Duration: {total_work_dur}", 
            f"Total Break Duration: {total_break_dur}"
        ])
        ws.append([])

        ws.append(headers)
        
        # Style headers
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        for cell in ws[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for r in data:
            ws.append([
                r['date'], r['employee_id'], r['employee_name'], r['email'],
                r['department'], r['project'], r['task'], r['description'], r['status'], 
                r['start_time'], r['end_time'],
                r['work_duration'], r['break_duration'], r['total_daily_duration']
            ])

        # Totals row (use numeric values for Excel formatting)
        total_row = [
            'TOTALS', '', '', '', '', '', '', '', '', '', '',
            total_work_seconds / 86400.0, 
            total_break_seconds / 86400.0, 
            total_daily_seconds_sum / 86400.0
        ]
        ws.append(total_row)
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        total_font = Font(name="Arial", size=10, bold=True, color="0F172A")
        for i, cell in enumerate(ws[ws.max_row], start=1):
            cell.fill = total_fill
            cell.font = total_font
            if i in [12, 13, 14]:
                cell.number_format = '[h]:mm:ss'
            
        # Adjust column widths
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 40), 12)
            
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        
        response = HttpResponse(
            stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        response['Access-Control-Expose-Headers'] = 'Content-Disposition, X-Suggested-Filename'
        response['X-Suggested-Filename'] = f"{filename}.xlsx"
        return response

    return Response({'error': 'Invalid format. Use csv or xlsx.'}, status=400)


# ==========================================
# SETTINGS
# ==========================================

@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def settings_view(request):
    if request.method == 'GET':
        return Response(storage.get_settings())
        
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)
        
    start_time = request.data.get('work_start_time')
    end_time = request.data.get('work_end_time')
    tz = request.data.get('timezone', 'Asia/Kolkata')
    
    updates = {}
    if start_time:
        updates['work_start_time'] = start_time
    if end_time:
        updates['work_end_time'] = end_time
    if tz:
        updates['timezone'] = tz
        
    storage.update_settings(updates, admin_email=request.user.email)
    return Response({'message': 'Settings updated successfully.', 'settings': storage.get_settings()})


# ==========================================
# MONTHLY TIMESHEET (EMPLOYEE & ADMIN)
# ==========================================

def get_monthly_timesheet_data(target_user, year, month, filter_project=None, filter_task=None):
    num_days = calendar.monthrange(year, month)[1]
    
    all_events = storage.get_events_for_user(target_user.id)
    timesheet_entries = storage.get_timesheet_entries_for_month(target_user.id, year, month)
    entries_by_date = {e['date']: e for e in timesheet_entries}
    
    settings = storage.get_settings()
    work_start_h, work_start_m = map(int, settings.get('work_start_time', '10:00').split(':'))
    work_end_h, work_end_m = map(int, settings.get('work_end_time', '19:00').split(':'))

    now_utc = datetime.now(timezone.utc)
    now_ist = now_utc.astimezone(IST_TZ)
    today_ist_date = now_ist.date()

    days_data = []
    total_work_seconds = 0
    total_break_seconds = 0
    days_worked_count = 0
    late_starts_count = 0
    missing_clockouts_count = 0

    sr_no = 1
    for day_num in range(1, num_days + 1):
        day_date = date(year, month, day_num)
        date_str = day_date.isoformat()
        day_of_week = day_date.strftime('%A')
        date_formatted = day_date.strftime('%d-%b-%Y')
        is_weekend = day_date.weekday() >= 5

        # Filter events for this day
        day_events = [
            e for e in all_events
            if parse_iso_utc(e['timestamp']) and parse_iso_utc(e['timestamp']).astimezone(IST_TZ).date() == day_date
        ]

        metrics = calculate_day_metrics(day_events, now_utc if day_date == today_ist_date else None)
        
        # Timesheet manual entry
        saved_entry = entries_by_date.get(date_str, {})
        project_name = saved_entry.get('project_name', '')
        task_name = saved_entry.get('task_name', '')
        work_description = saved_entry.get('work_description', '')

        # Filter check if searching by project or task
        if filter_project and filter_project.lower() not in project_name.lower():
            continue
        if filter_task and filter_task.lower() not in task_name.lower():
            continue

        # Determine status
        is_late = False
        is_missing = False

        if day_events:
            start_dt = parse_iso_utc(metrics['start_time']).astimezone(IST_TZ)
            office_start_dt = start_dt.replace(hour=work_start_h, minute=work_start_m, second=0, microsecond=0)
            if start_dt > office_start_dt + timedelta(minutes=5):
                is_late = True
                late_starts_count += 1

            if metrics['end_time']:
                day_status = 'Completed'
                days_worked_count += 1
            else:
                if day_date == today_ist_date:
                    day_status = 'In Progress'
                    days_worked_count += 1
                else:
                    day_status = 'Missing Clock-out'
                    is_missing = True
                    missing_clockouts_count += 1
                    days_worked_count += 1
        else:
            if is_weekend:
                day_status = 'Weekend'
            elif day_date > today_ist_date:
                day_status = 'Upcoming'
            else:
                day_status = 'No Record'

        total_work_seconds += metrics['total_work_seconds']
        total_break_seconds += metrics['total_break_seconds']

        start_time_str = '-'
        if metrics['start_time']:
            start_time_str = parse_iso_utc(metrics['start_time']).astimezone(IST_TZ).strftime('%I:%M:%S %p')

        end_time_str = '-'
        if metrics['end_time']:
            end_time_str = parse_iso_utc(metrics['end_time']).astimezone(IST_TZ).strftime('%I:%M:%S %p')

        work_hours = round(metrics['total_work_seconds'] / 3600.0, 2)
        break_hours = round(metrics['total_break_seconds'] / 3600.0, 2)
        
        # Formatted duration string hh:mm:ss
        b_h = metrics['total_break_seconds'] // 3600
        b_m = (metrics['total_break_seconds'] % 3600) // 60
        b_s = metrics['total_break_seconds'] % 60
        break_duration_str = f"{b_h:02d}:{b_m:02d}:{b_s:02d}"

        w_h = metrics['total_work_seconds'] // 3600
        w_m = (metrics['total_work_seconds'] % 3600) // 60
        w_s = metrics['total_work_seconds'] % 60
        work_duration_str = f"{w_h:02d}:{w_m:02d}:{w_s:02d}"

        total_day_seconds = 0
        now_utc = datetime.now(timezone.utc)
        if metrics['start_time']:
            s_dt = parse_iso_utc(metrics['start_time'])
            if metrics['end_time']:
                e_dt = parse_iso_utc(metrics['end_time'])
                total_day_seconds = max(0, int((e_dt - s_dt).total_seconds()))
            elif day_status != 'Not Started':
                total_day_seconds = max(0, int((now_utc - s_dt).total_seconds()))

        required_seconds = 32400  # 9 hours
        is_completed = total_day_seconds >= required_seconds
        completion_status = 'Completed' if is_completed else 'Not Completed'
        diff_sec = max(0, required_seconds - total_day_seconds) if not is_completed else max(0, total_day_seconds - required_seconds)

        td_h = total_day_seconds // 3600
        td_m = (total_day_seconds % 3600) // 60
        td_s = total_day_seconds % 60
        total_day_duration_str = f"{td_h:02d}:{td_m:02d}:{td_s:02d}"

        req_h = required_seconds // 3600
        required_hours_str = f"{req_h:02d}:00:00"

        diff_h = diff_sec // 3600
        diff_m = (diff_sec % 3600) // 60
        diff_s = diff_sec % 60
        diff_str = f"{diff_h:02d}:{diff_m:02d}:{diff_s:02d}"

        days_data.append({
            'sr_no': sr_no,
            'date': date_str,
            'date_formatted': date_formatted,
            'day_of_week': day_of_week,
            'start_time': start_time_str,
            'end_time': end_time_str,
            'start_time_iso': metrics['start_time'],
            'end_time_iso': metrics['end_time'],
            'total_day_seconds': total_day_seconds,
            'total_day_duration_str': total_day_duration_str,
            'total_work_seconds': metrics['total_work_seconds'],
            'work_hours': work_hours,
            'work_duration_str': work_duration_str,
            'total_break_seconds': metrics['total_break_seconds'],
            'break_hours': break_hours,
            'break_duration_str': break_duration_str,
            'required_hours_str': required_hours_str,
            'completion_status': completion_status,
            'diff_duration_str': diff_str,
            'project_name': project_name,
            'task_name': task_name,
            'work_description': work_description,
            'status': day_status,
            'is_late_start': is_late,
            'is_missing_clockout': is_missing,
            'is_weekend': is_weekend,
            'updated_at': saved_entry.get('updated_at', '')
        })
        sr_no += 1

    summary = {
        'total_working_hours': round(total_work_seconds / 3600.0, 2),
        'total_break_hours': round(total_break_seconds / 3600.0, 2),
        'total_work_seconds': total_work_seconds,
        'total_break_seconds': total_break_seconds,
        'days_worked': days_worked_count,
        'late_starts': late_starts_count,
        'missing_clockouts': missing_clockouts_count
    }

    return {
        'employee': {
            'id': target_user.id,
            'employee_id': target_user.employee_id or f"EMP{target_user.id:03d}",
            'name': f"{target_user.first_name} {target_user.last_name}".strip() or target_user.username,
            'email': target_user.email,
            'department': target_user.department or 'Unassigned',
            'role': target_user.role
        },
        'year': year,
        'month': month,
        'month_name': calendar.month_name[month],
        'days': days_data,
        'summary': summary
    }


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def monthly_timesheet_view(request):
    now_ist = datetime.now(timezone.utc).astimezone(IST_TZ)
    try:
        year = int(request.GET.get('year', now_ist.year))
        month = int(request.GET.get('month', now_ist.month))
    except (ValueError, TypeError):
        year = now_ist.year
        month = now_ist.month

    if month < 1 or month > 12:
        return Response({'error': 'Month must be between 1 and 12.'}, status=400)

    requested_user_id = request.GET.get('user_id')
    
    # Permission verification: Employee cannot view other users' timesheets
    if request.user.role != 'ADMIN':
        if requested_user_id and int(requested_user_id) != request.user.id:
            return Response({'error': 'Access denied. Employees can view only their own timesheets.'}, status=403)
        target_user = request.user
    else:
        if requested_user_id:
            try:
                target_user = User.objects.get(id=int(requested_user_id), role='EMPLOYEE')
            except User.DoesNotExist:
                return Response({'error': 'Employee not found.'}, status=404)
        else:
            first_emp = User.objects.filter(role='EMPLOYEE', is_active=True).first()
            if not first_emp:
                first_emp = User.objects.filter(role='EMPLOYEE').first()
            if first_emp:
                target_user = first_emp
            else:
                return Response({'error': 'No employees found in the system.'}, status=404)

    filter_project = request.GET.get('project')
    filter_task = request.GET.get('task')

    data = get_monthly_timesheet_data(target_user, year, month, filter_project, filter_task)
    return Response(data)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def save_timesheet_entry_view(request):
    work_date = request.data.get('date')
    project_name = (request.data.get('project_name') or '').strip()
    task_name = (request.data.get('task_name') or '').strip()
    work_description = (request.data.get('work_description') or '').strip()
    
    if not work_date:
        return Response({'error': 'Date is required (YYYY-MM-DD).'}, status=400)

    try:
        datetime.strptime(work_date, '%Y-%m-%d')
    except ValueError:
        return Response({'error': 'Invalid date format. Expected YYYY-MM-DD.'}, status=400)

    requested_user_id = request.data.get('user_id')

    # Permission check: Employee can edit only their own timesheet entries
    if request.user.role != 'ADMIN':
        if requested_user_id and int(requested_user_id) != request.user.id:
            return Response({'error': 'Access denied. Employees can only update their own timesheet entries.'}, status=403)
        target_user_id = request.user.id
    else:
        if requested_user_id:
            if not User.objects.filter(id=int(requested_user_id), role='EMPLOYEE').exists():
                return Response({'error': 'Target user is not an employee.'}, status=400)
            target_user_id = int(requested_user_id)
        else:
            first_emp = User.objects.filter(role='EMPLOYEE').first()
            if not first_emp:
                return Response({'error': 'No employee found to save timesheet entry.'}, status=400)
            target_user_id = first_emp.id

    entry = storage.save_timesheet_entry(
        user_id=target_user_id,
        work_date=work_date,
        project_name=project_name,
        task_name=task_name,
        work_description=work_description,
        updated_by=request.user.email
    )
    return Response({'message': 'Timesheet entry saved successfully.', 'entry': entry})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_timesheet_view(request):
    now_ist = datetime.now(timezone.utc).astimezone(IST_TZ)
    try:
        year = int(request.GET.get('year', now_ist.year))
        month = int(request.GET.get('month', now_ist.month))
    except (ValueError, TypeError):
        year = now_ist.year
        month = now_ist.month

    requested_user_id = request.GET.get('user_id')
    if request.user.role != 'ADMIN':
        if requested_user_id and int(requested_user_id) != request.user.id:
            return Response({'error': 'Access denied. Employees can only export their own timesheets.'}, status=403)
        target_user = request.user
    else:
        if requested_user_id:
            try:
                target_user = User.objects.get(id=int(requested_user_id), role='EMPLOYEE')
            except User.DoesNotExist:
                return Response({'error': 'Employee not found.'}, status=404)
        else:
            first_emp = User.objects.filter(role='EMPLOYEE', is_active=True).first()
            if not first_emp:
                first_emp = User.objects.filter(role='EMPLOYEE').first()
            if first_emp:
                target_user = first_emp
            else:
                return Response({'error': 'No employees found in the system.'}, status=404)

    filter_project = request.GET.get('project')
    filter_task = request.GET.get('task')
    export_format = request.GET.get('format', 'csv').lower()

    data = get_monthly_timesheet_data(target_user, year, month, filter_project, filter_task)
    emp = data['employee']
    summary = data['summary']
    days = data['days']

    storage.log_audit(
        action="EXPORT_TIMESHEET",
        actor=request.user.email,
        details={"target_user_id": emp['id'], "year": year, "month": month, "format": export_format}
    )

    filename_prefix = f"timesheet_{emp['name'].replace(' ', '_')}_{year}_{month:02d}"

    if export_format == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename_prefix}.csv"'
        
        writer = csv.writer(response)
        # Summary Header (based on Reference Image 2)
        writer.writerow(['KNOWLETIVE MONTHLY EMPLOYEE TIMESHEET'])
        writer.writerow(['Employee ID:', emp.get('employee_id', ''), 'Employee:', emp['name'], 'Email:', emp['email'], 'Department:', emp['department']])
        writer.writerow(['Month & Year:', f"{data['month_name']} {year}"])
        writer.writerow(['Total Working Hrs:', f"{summary['total_working_hours']} hrs", 'Total Break Time:', f"{summary['total_break_hours']} hrs", 'Days Worked:', summary['days_worked']])
        writer.writerow([]) # blank line

        # Table Header (based on Reference Image 1)
        writer.writerow([
            'Sr. No.', 'Date (IST)', 'Day', 'Project', 'Task Name', 'Description', 
            'Clock In', 'Clock Out', 'Total Day Time', 'Actual Work Time', 'Break Time',
            'Required Hours', 'Completed Hours Status', 'Shortfall / Extra Time', 'Status'
        ])
        for r in days:
            writer.writerow([
                r['sr_no'], r['date_formatted'], r['day_of_week'], r['project_name'], 
                r['task_name'], r['work_description'], r['start_time'], r['end_time'],
                r['total_day_duration_str'], r['work_duration_str'], r['break_duration_str'],
                r['required_hours_str'], r['completion_status'], r['diff_duration_str'], r['status']
            ])
        return response

    elif export_format in ['xlsx', 'excel']:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{data['month_name'][:3]} {year} Timesheet"

        # Summary Block (Reference Image 2 styling)
        title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        ws.append(['KNOWLETIVE MONTHLY EMPLOYEE TIMESHEET'])
        ws['A1'].font = title_font

        ws.append([f"Employee ID: {emp.get('employee_id', '')}", f"Employee: {emp['name']}", f"Email: {emp['email']}", f"Department: {emp['department']}"])
        ws.append([f"Period: {data['month_name']} {year}"])
        ws.append([])

        summary_header = ['Total Break Time', 'Total Working Hrs.', 'Productive Days']
        summary_values = [f"{summary['total_break_hours']} hrs", f"{summary['total_working_hours']} hrs", f"{summary['days_worked']} days"]
        
        ws.append(summary_header)
        ws.append(summary_values)
        
        # Color summary cards
        header_fill_pink = PatternFill(start_color="FCE7F3", end_color="FCE7F3", fill_type="solid")
        for col_idx in range(1, 4):
            cell = ws.cell(row=5, column=col_idx)
            cell.fill = header_fill_pink
            cell.font = Font(name="Arial", size=10, bold=True, color="9D174D")
            cell.alignment = Alignment(horizontal="center")
            
            val_cell = ws.cell(row=6, column=col_idx)
            val_cell.font = Font(name="Arial", size=12, bold=True, color="111827")
            val_cell.alignment = Alignment(horizontal="center")

        ws.append([])

        # Table Header (Reference Image 1 pink/dark table header)
        table_headers = [
            'Sr. No.', 'Project', 'Task Name', 'Description', 'Clock In', 'Clock Out', 
            'Total Day Time', 'Actual Work Time', 'Break Time', 'Required Hours', 'Completed Hours Status', 'Shortfall / Extra Time', 'Status'
        ]
        ws.append(table_headers)
        
        table_header_fill = PatternFill(start_color="FBCFE8", end_color="FBCFE8", fill_type="solid")
        for cell in ws[8]:
            cell.fill = table_header_fill
            cell.font = Font(name="Arial", size=11, bold=True, color="831843")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r in days:
            ws.append([
                r['sr_no'],
                r['project_name'] or '-',
                r['task_name'] or '-',
                r['work_description'] or '-',
                f"{r['date_formatted']} {r['start_time']}" if r['start_time'] != '-' else '-',
                f"{r['date_formatted']} {r['end_time']}" if r['end_time'] != '-' else '-',
                r['total_day_duration_str'],
                r['work_duration_str'],
                r['break_duration_str'],
                r['required_hours_str'],
                r['completion_status'],
                r['diff_duration_str'],
                r['status']
            ])

        # Formatting widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        response = HttpResponse(
            stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename_prefix}.xlsx"'
        return response

    return Response({'error': 'Invalid format. Use csv or xlsx.'}, status=400)


# ==========================================
# TIME-WISE TASK & PROJECT LOGS
# ==========================================

def format_task_seconds(sec):
    h = sec // 3600
    m = (sec % 3600) // 60
    s = sec % 60
    
    if h > 0:
        if m > 0 and s > 0:
            return f"{h} hr {m} min {s} sec"
        elif m > 0:
            return f"{h} hr {m} min"
        elif s > 0:
            return f"{h} hr {s} sec"
        else:
            return f"{h} hr"
    elif m > 0:
        if s > 0:
            return f"{m} min {s} sec"
        else:
            return f"{m} min"
    else:
        return f"{s} sec"


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def daily_tasks_view(request):
    now_utc = datetime.now(timezone.utc)
    now_ist = now_utc.astimezone(IST_TZ)
    target_date = request.GET.get('date', now_ist.date().isoformat())

    requested_user_id = request.GET.get('user_id')
    if request.user.role != 'ADMIN':
        if requested_user_id and int(requested_user_id) != request.user.id:
            return Response({'error': 'Access denied. You can only view your own task log.'}, status=403)
        target_user = request.user
    else:
        if requested_user_id:
            try:
                target_user = User.objects.get(id=int(requested_user_id), role='EMPLOYEE')
            except User.DoesNotExist:
                return Response({'error': 'Employee not found.'}, status=404)
        else:
            first_emp = User.objects.filter(role='EMPLOYEE', is_active=True).first()
            if not first_emp:
                first_emp = User.objects.filter(role='EMPLOYEE').first()
            if first_emp:
                target_user = first_emp
            else:
                return Response({'error': 'No employees found in the system.'}, status=404)

    filter_project = request.GET.get('project')
    raw_tasks = storage.get_tasks_for_user_and_date(target_user.id, target_date)
    if filter_project and filter_project.strip():
        raw_tasks = [t for t in raw_tasks if (t.get('project_name') or '').strip().lower() == filter_project.strip().lower()]

    processed_tasks = []
    total_task_seconds = 0
    project_totals = {}
    active_task = None

    for idx, t in enumerate(raw_tasks, start=1):
        start_dt = parse_iso_utc(t['start_time'])
        end_dt = parse_iso_utc(t['end_time']) if t['end_time'] else None

        dur_sec = storage.calculate_task_duration(t, now_utc)
        total_task_seconds += dur_sec

        proj = t['project_name'] or 'Unassigned'
        if proj not in project_totals:
            project_totals[proj] = {'total_seconds': 0, 'count': 0}
        project_totals[proj]['total_seconds'] += dur_sec
        project_totals[proj]['count'] += 1

        start_time_ist = start_dt.astimezone(IST_TZ).strftime('%I:%M:%S %p') if start_dt else '-'
        end_time_ist = end_dt.astimezone(IST_TZ).strftime('%I:%M:%S %p') if end_dt else ('Paused (Break)' if t.get('status') == 'PAUSED' else 'In progress')

        task_item = {
            'sr_no': idx,
            'task_log_id': t['task_log_id'],
            'employee_id': t.get('employee_id') or target_user.employee_id or f"EMP{target_user.id:03d}",
            'employee_name': t.get('employee_name') or f"{target_user.first_name} {target_user.last_name}".strip() or target_user.username,
            'project_name': t['project_name'],
            'task_name': t['task_name'],
            'description': t['description'],
            'start_time': t['start_time'],
            'end_time': t['end_time'],
            'start_time_formatted': start_time_ist,
            'end_time_formatted': end_time_ist,
            'duration_seconds': dur_sec,
            'duration_seconds_base': int(t.get('duration_seconds') or 0),
            'current_segment_start': t.get('current_segment_start'),
            'work_segments': t.get('work_segments', []),
            'duration_hours': round(dur_sec / 3600.0, 2),
            'duration_formatted': format_task_seconds(dur_sec),
            'status': t['status']
        }
        processed_tasks.append(task_item)

        if t['status'] in ['IN_PROGRESS', 'PAUSED']:
            active_task = task_item

    project_breakdown = [
        {
            'project_name': p_name,
            'total_seconds': data['total_seconds'],
            'total_hours': round(data['total_seconds'] / 3600.0, 2),
            'duration_formatted': format_task_seconds(data['total_seconds']),
            'count': data['count']
        }
        for p_name, data in sorted(project_totals.items(), key=lambda x: x[1]['total_seconds'], reverse=True)
    ]

    active_count = sum(1 for t in processed_tasks if t['status'] == 'IN_PROGRESS')
    completed_count = sum(1 for t in processed_tasks if t['status'] != 'IN_PROGRESS')
    emp_code = target_user.employee_id or f"EMP{target_user.id:03d}"
    emp_name = f"{target_user.first_name} {target_user.last_name}".strip() or target_user.username

    summary = {
        'total_tasks': len(processed_tasks),
        'active_in_progress': active_count,
        'completed': completed_count,
        'total_seconds': total_task_seconds,
        'total_hours': round(total_task_seconds / 3600.0, 2)
    }

    by_employee = [{
        'user_id': target_user.id,
        'employee_id': emp_code,
        'employee_name': emp_name,
        'email': target_user.email,
        'department': target_user.department or 'General',
        'total_seconds': total_task_seconds,
        'total_hours': round(total_task_seconds / 3600.0, 2),
        'duration_formatted': format_task_seconds(total_task_seconds),
        'tasks_count': len(processed_tasks)
    }] if processed_tasks else []

    by_project_and_employee = [
        {
            'project_name': p['project_name'],
            'employee_name': emp_name,
            'total_seconds': p['total_seconds'],
            'total_hours': p['total_hours'],
            'duration_formatted': p['duration_formatted'],
            'tasks_count': p['count']
        }
        for p in project_breakdown
    ]

    return Response({
        'date': target_date,
        'employee': {
            'id': target_user.id,
            'employee_id': emp_code,
            'name': emp_name,
            'email': target_user.email,
            'department': target_user.department
        },
        'tasks': processed_tasks,
        'active_task': active_task,
        'total_task_seconds': total_task_seconds,
        'total_task_hours': round(total_task_seconds / 3600.0, 2),
        'total_task_formatted': format_task_seconds(total_task_seconds),
        'project_breakdown': project_breakdown,
        'summary': summary,
        'by_employee': by_employee,
        'by_project_and_employee': by_project_and_employee
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def start_task_view(request):
    if request.user.role != 'EMPLOYEE':
        return Response({'error': 'Only employee accounts can log tasks.'}, status=403)

    project_name = (request.data.get('project_name') or '').strip()
    task_name = (request.data.get('task_name') or '').strip()
    description = (request.data.get('description') or '').strip()

    if not project_name:
        return Response({'error': 'Project name is required to start a task.'}, status=400)
    if not task_name:
        return Response({'error': 'Task name is required to start a task.'}, status=400)

    now_utc = datetime.now(timezone.utc)
    now_ist = now_utc.astimezone(IST_TZ)
    today_ist_date = now_ist.date()

    events = storage.get_events_for_user(request.user.id)
    today_events = [
        e for e in events 
        if parse_iso_utc(e['timestamp']) and parse_iso_utc(e['timestamp']).astimezone(IST_TZ).date() == today_ist_date
    ]
    today_events.sort(key=lambda x: x['timestamp'])

    metrics = calculate_day_metrics(today_events, now_utc)
    current_status = metrics['status']

    if current_status == 'Ended':
        return Response({'error': 'Workday has already ended for today. Contact admin for corrections.'}, status=400)

    # CASE 4: USER ADDS / STARTS A NEW TASK DURING BREAK
    # If employee is currently on BREAK and adds/starts a new task:
    # 1. Close current break session -> automatically record RESUME_WORK
    # 2. Transition state to WORKING
    # 3. New task timer starts, break timer stops
    if current_status == 'On Break':
        storage.record_event(request.user.id, 'RESUME_WORK', source='web')
    elif current_status == 'Not Started':
        storage.record_event(request.user.id, 'START_WORK', source='web')

    # Attach permanent Employee ID and saved employee name from authenticated profile
    emp_id = getattr(request.user, 'employee_id', '') or f"EMP{request.user.id:03d}"
    emp_name = f"{request.user.first_name} {request.user.last_name}".strip() or request.user.username

    try:
        new_task = storage.start_task(
            user_id=request.user.id,
            project_name=project_name,
            task_name=task_name,
            description=description,
            user_email=request.user.email,
            employee_id=emp_id,
            employee_name=emp_name
        )
        return Response({'message': 'Task started successfully.', 'task': new_task}, status=201)
    except ValueError as e:
        return Response({'error': str(e)}, status=400)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def end_task_view(request):
    if request.user.role != 'EMPLOYEE':
        return Response({'error': 'Only employee accounts can log tasks.'}, status=403)

    task_log_id = request.data.get('task_log_id')
    try:
        ended_task = storage.end_task(
            user_id=request.user.id,
            task_log_id=task_log_id,
            user_email=request.user.email
        )
        return Response({'message': 'Task ended successfully.', 'task': ended_task})
    except ValueError as e:
        return Response({'error': str(e)}, status=400)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def update_task_view(request, task_log_id):
    if request.user.role != 'EMPLOYEE':
        return Response({'error': 'Only employee accounts can update tasks.'}, status=403)

    project_name = request.data.get('project_name')
    task_name = request.data.get('task_name')
    description = request.data.get('description')

    try:
        updated = storage.update_task_details(
            user_id=request.user.id,
            task_log_id=str(task_log_id),
            project_name=project_name,
            task_name=task_name,
            description=description,
            user_email=request.user.email
        )
        return Response({'message': 'Task details updated.', 'task': updated})
    except ValueError as e:
        return Response({'error': str(e)}, status=400)


def get_admin_task_report_data(request):
    now_ist = datetime.now(timezone.utc).astimezone(IST_TZ)
    target_date = request.GET.get('date', now_ist.date().isoformat())
    employee_id = request.GET.get('employee_id')
    project = request.GET.get('project')

    employee_users = User.objects.filter(role='EMPLOYEE')
    employee_ids = set(employee_users.values_list('id', flat=True))
    users_map = {u.id: u for u in employee_users}

    if employee_id:
        matched_users = employee_users
        if str(employee_id).isdigit():
            matched_users = matched_users.filter(models.Q(id=int(employee_id)) | models.Q(employee_id=employee_id))
        else:
            matched_users = matched_users.filter(employee_id=employee_id)
        matched_emp = matched_users.first()
        target_uid = matched_emp.id if matched_emp else None
        if not target_uid:
            all_tasks = []
        else:
            all_tasks = storage.get_all_tasks_for_date(
                date_str=target_date,
                filter_user_id=target_uid,
                filter_project=project
            )
            all_tasks = [t for t in all_tasks if t.get('user_id') in employee_ids]
    else:
        all_tasks = storage.get_all_tasks_for_date(
            date_str=target_date,
            filter_user_id=None,
            filter_project=project
        )
        all_tasks = [t for t in all_tasks if t.get('user_id') in employee_ids]

    now_utc = datetime.now(timezone.utc)
    task_rows = []
    emp_aggregates = {}
    proj_emp_aggregates = {}
    active_count = 0
    completed_count = 0

    for idx, t in enumerate(all_tasks, start=1):
        u = users_map.get(t['user_id'])
        emp_code = u.employee_id if (u and u.employee_id) else (t.get('employee_id') or (f"EMP{u.id:03d}" if u else f"EMP{t['user_id']:03d}"))
        emp_name = f"{u.first_name} {u.last_name}".strip() if u else (t.get('employee_name') or f"User {t['user_id']}")
        emp_email = u.email if u else '-'
        emp_dept = u.department if u else '-'

        dur_sec = storage.calculate_task_duration(t, now_utc)
        start_dt = parse_iso_utc(t.get('start_time'))
        end_dt = parse_iso_utc(t.get('end_time')) if t.get('end_time') else None
        if t.get('status') == 'IN_PROGRESS':
            active_count += 1
        else:
            completed_count += 1

        # Group by employee
        if t['user_id'] not in emp_aggregates:
            emp_aggregates[t['user_id']] = {
                'user_id': t['user_id'],
                'employee_id': emp_code,
                'employee_name': emp_name,
                'email': emp_email,
                'department': emp_dept,
                'total_seconds': 0,
                'tasks_count': 0
            }
        emp_aggregates[t['user_id']]['total_seconds'] += dur_sec
        emp_aggregates[t['user_id']]['tasks_count'] += 1

        # Group by project & employee
        proj_key = f"{t['project_name']}::{emp_name}"
        if proj_key not in proj_emp_aggregates:
            proj_emp_aggregates[proj_key] = {
                'project_name': t['project_name'],
                'employee_name': emp_name,
                'total_seconds': 0,
                'tasks_count': 0
            }
        proj_emp_aggregates[proj_key]['total_seconds'] += dur_sec
        proj_emp_aggregates[proj_key]['tasks_count'] += 1

        start_str = start_dt.astimezone(IST_TZ).strftime('%I:%M:%S %p') if start_dt else '-'
        end_str = end_dt.astimezone(IST_TZ).strftime('%I:%M:%S %p') if end_dt else 'In progress'

        task_rows.append({
            'sr_no': idx,
            'task_log_id': t['task_log_id'],
            'user_id': t['user_id'],
            'employee_id': emp_code,
            'employee_name': emp_name,
            'email': emp_email,
            'department': emp_dept,
            'project_name': t['project_name'],
            'task_name': t['task_name'],
            'description': t['description'],
            'start_time': start_str,
            'end_time': end_str,
            'duration_seconds': dur_sec,
            'duration_hours': round(dur_sec / 3600.0, 2),
            'duration_formatted': format_task_seconds(dur_sec),
            'status': t['status']
        })

    by_employee = [
        {
            'user_id': d['user_id'],
            'employee_id': d['employee_id'],
            'employee_name': d['employee_name'],
            'email': d['email'],
            'department': d['department'],
            'total_seconds': d['total_seconds'],
            'total_hours': round(d['total_seconds'] / 3600.0, 2),
            'duration_formatted': format_task_seconds(d['total_seconds']),
            'tasks_count': d['tasks_count']
        }
        for d in sorted(emp_aggregates.values(), key=lambda x: x['total_seconds'], reverse=True)
    ]

    by_project_and_employee = [
        {
            'project_name': d['project_name'],
            'employee_name': d['employee_name'],
            'total_seconds': d['total_seconds'],
            'total_hours': round(d['total_seconds'] / 3600.0, 2),
            'duration_formatted': format_task_seconds(d['total_seconds']),
            'tasks_count': d['tasks_count']
        }
        for d in sorted(proj_emp_aggregates.values(), key=lambda x: x['total_seconds'], reverse=True)
    ]

    return {
        'date': target_date,
        'tasks': task_rows,
        'by_employee': by_employee,
        'by_project_and_employee': by_project_and_employee,
        'summary': {
            'total_tasks': len(task_rows),
            'active_in_progress': active_count,
            'completed': completed_count,
            'total_seconds': sum(r['duration_seconds'] for r in task_rows),
            'total_hours': round(sum(r['duration_seconds'] for r in task_rows) / 3600.0, 2)
        }
    }


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_tasks_report_view(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)

    data = get_admin_task_report_data(request)
    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_tasks_export_view(request):
    if request.user.role != 'ADMIN':
        return Response({'error': 'Admin permissions required.'}, status=403)

    export_format = request.GET.get('format', 'csv').lower()
    data = get_admin_task_report_data(request)
    date_str = data['date']
    tasks = data['tasks']
    summary = data['summary']

    storage.log_audit(
        action="EXPORT_TASK_REPORT",
        actor=request.user.email,
        details={"date": date_str, "format": export_format, "tasks_count": len(tasks)}
    )

    filename_prefix = f"daily_tasks_report_{date_str}"

    if export_format == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename_prefix}.csv"'

        writer = csv.writer(response)
        writer.writerow(['KNOWLETIVE DAILY TIME-WISE TASK REPORT', f"Date: {date_str}"])
        writer.writerow([
            'Total Tasks:', summary['total_tasks'], 
            'Active in Progress:', summary['active_in_progress'], 
            'Total Time:', format_task_seconds(summary['total_seconds'])
        ])
        writer.writerow([])
        writer.writerow([
            'Sr. No.', 'Employee ID', 'Employee Name', 'Department', 'Project Name', 
            'Task Name', 'Work Description', 'Start Time (IST)', 'End Time (IST)', 
            'Duration (Hours)', 'Duration', 'Status'
        ])
        for t in tasks:
            writer.writerow([
                t['sr_no'], t['employee_id'], t['employee_name'], t['department'], t['project_name'],
                t['task_name'], t['description'], t['start_time'], t['end_time'],
                t['duration_hours'], t['duration_formatted'], t['status']
            ])
        return response

    elif export_format in ['xlsx', 'excel']:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Tasks {date_str}"

        # Header block
        title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        ws.append(['KNOWLETIVE DAILY TIME-WISE TASK REPORT'])
        ws['A1'].font = title_font
        ws.append([f"Date (IST): {date_str}", f"Total Tasks: {summary['total_tasks']}", f"Total Time: {format_task_seconds(summary['total_seconds'])}"])
        ws.append([])

        headers = [
            'Sr. No.', 'Employee ID', 'Employee Name', 'Department', 'Project Name', 
            'Task Name', 'Work Description', 'Start Time (IST)', 'End Time (IST)', 
            'Duration (Hours)', 'Duration', 'Status'
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        for cell in ws[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for t in tasks:
            ws.append([
                t['sr_no'], t['employee_id'], t['employee_name'], t['department'], t['project_name'],
                t['task_name'], t['description'], t['start_time'], t['end_time'],
                t['duration_hours'], t['duration_formatted'], t['status']
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        response = HttpResponse(
            stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename_prefix}.xlsx"'
        return response

    return Response({'error': 'Invalid format. Use csv or xlsx.'}, status=400)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def analytics_dashboard_view(request):
    """
    Role-enforced analytics dashboard API.
    - Employees: strictly limited to their own data (request.user.id).
    - Admins: can view all permitted employees or filter by a specific employee.
    Calculates summary metrics, daily trends, project breakdowns, employee breakdowns, and task breakdowns using real data.
    """
    user = request.user
    is_admin = (user.role == 'ADMIN')

    # 1. Scope users based on role
    employee_users = User.objects.filter(role='EMPLOYEE')
    users_map = {u.id: u for u in employee_users}
    users_map[user.id] = user

    if not is_admin:
        # Strictly scope employees to their own records only; ignore any other employee_id parameter
        target_users = [user]
        target_user_ids = {user.id}
    else:
        req_emp = request.GET.get('employee_id')
        if req_emp:
            if str(req_emp).isdigit():
                matched = employee_users.filter(models.Q(id=int(req_emp)) | models.Q(employee_id=req_emp)).first()
            else:
                matched = employee_users.filter(employee_id=req_emp).first()
            target_users = [matched] if matched else []
            target_user_ids = {matched.id} if matched else set()
        else:
            target_users = list(employee_users)
            target_user_ids = set(employee_users.values_list('id', flat=True))

    # 2. Date filters
    now_utc = datetime.now(timezone.utc)
    now_ist = now_utc.astimezone(IST_TZ)
    today_ist = now_ist.date()

    start_date_param = request.GET.get('start_date')
    end_date_param = request.GET.get('end_date')
    project_filter = request.GET.get('project', '').strip()
    period = request.GET.get('period', '7d')

    if start_date_param and end_date_param:
        try:
            d_start = date.fromisoformat(start_date_param)
            d_end = date.fromisoformat(end_date_param)
        except ValueError:
            d_end = today_ist
            d_start = today_ist - timedelta(days=6)
    elif period == 'today':
        d_start = today_ist
        d_end = today_ist
    elif period == '30d':
        d_end = today_ist
        d_start = today_ist - timedelta(days=29)
    elif period == 'month':
        d_start = today_ist.replace(day=1)
        days_in_m = calendar.monthrange(today_ist.year, today_ist.month)[1]
        d_end = today_ist.replace(day=days_in_m)
    else:  # default 7d
        d_end = today_ist
        d_start = today_ist - timedelta(days=6)

    if d_start > d_end:
        d_start, d_end = d_end, d_start

    d_start_str = d_start.isoformat()
    d_end_str = d_end.isoformat()

    # 3. Retrieve all real events and tasks in date range for target users
    all_events = storage.get_all_events() if is_admin else storage.get_events_for_user(user.id)
    scoped_events = [
        e for e in all_events
        if int(e.get('user_id', -1)) in target_user_ids
    ]

    filter_uid = user.id if not is_admin else (list(target_user_ids)[0] if len(target_user_ids) == 1 and request.GET.get('employee_id') else None)
    all_tasks = storage.get_all_tasks(
        start_date=d_start_str,
        end_date=d_end_str,
        user_id=filter_uid,
        project=(project_filter if project_filter else None)
    )
    scoped_tasks = [t for t in all_tasks if int(t.get('user_id', -1)) in target_user_ids]

    # 4. Daily time-tracking metrics aggregation
    daily_trend = []
    total_work_seconds = 0
    total_break_seconds = 0
    days_worked_set = set()
    active_employees_set = set()

    events_by_user_and_date = {}
    for ev in scoped_events:
        dt = parse_iso_utc(ev.get('timestamp'))
        if dt:
            ev_date = dt.astimezone(IST_TZ).date()
            if d_start <= ev_date <= d_end:
                u_id = int(ev.get('user_id', -1))
                events_by_user_and_date.setdefault((u_id, ev_date), []).append(ev)

    tasks_by_date = {}
    for t in scoped_tasks:
        t_date = t.get('date', '')
        tasks_by_date.setdefault(t_date, []).append(t)

    cur_day = d_start
    while cur_day <= d_end:
        cur_day_str = cur_day.isoformat()
        day_work_sec = 0
        day_break_sec = 0
        day_active_users = set()

        for u in target_users:
            u_day_events = events_by_user_and_date.get((u.id, cur_day), [])
            if u_day_events:
                metrics = calculate_day_metrics(u_day_events, now_utc if cur_day == today_ist else None)
                day_work_sec += metrics['total_work_seconds']
                day_break_sec += metrics['total_break_seconds']
                if metrics['total_work_seconds'] > 0:
                    day_active_users.add(u.id)
                    days_worked_set.add(cur_day_str)
                    active_employees_set.add(u.id)

        day_tasks = tasks_by_date.get(cur_day_str, [])
        for t in day_tasks:
            t_dur = storage.calculate_task_duration(t, now_utc)
            active_employees_set.add(int(t.get('user_id', -1)))
            days_worked_set.add(cur_day_str)
            if not day_work_sec and t_dur:
                day_work_sec += t_dur

        total_work_seconds += day_work_sec
        total_break_seconds += day_break_sec

        formatted_label = cur_day.strftime('%d %b')
        day_name = cur_day.strftime('%a')

        daily_trend.append({
            'date': cur_day_str,
            'label': formatted_label,
            'day_name': day_name,
            'work_hours': round(day_work_sec / 3600.0, 2),
            'work_seconds': day_work_sec,
            'break_hours': round(day_break_sec / 3600.0, 2),
            'break_seconds': day_break_sec,
            'tasks_count': len(day_tasks),
            'active_employees': len(day_active_users) if day_active_users else (len(set(t['user_id'] for t in day_tasks)) if day_tasks else 0)
        })

        cur_day += timedelta(days=1)

    # 5. Project & Task Aggregations
    project_aggregates = {}
    task_aggregates = {}
    completed_tasks_count = 0
    in_progress_tasks_count = 0

    for t in scoped_tasks:
        p_name = (t.get('project_name') or 'Unassigned').strip()
        t_name = (t.get('task_name') or 'Unnamed Task').strip()
        status = t.get('status', 'COMPLETED')

        s_dt = parse_iso_utc(t.get('start_time'))
        e_dt = parse_iso_utc(t.get('end_time')) if t.get('end_time') else None
        if e_dt and s_dt:
            dur = max(0, int((e_dt - s_dt).total_seconds()))
            completed_tasks_count += 1
        elif s_dt:
            dur = max(0, int((now_utc - s_dt).total_seconds()))
            in_progress_tasks_count += 1
        else:
            dur = 0

        if p_name not in project_aggregates:
            project_aggregates[p_name] = {
                'project_name': p_name,
                'total_seconds': 0,
                'tasks_count': 0,
                'employees': set()
            }
        project_aggregates[p_name]['total_seconds'] += dur
        project_aggregates[p_name]['tasks_count'] += 1
        if t.get('user_id'):
            project_aggregates[p_name]['employees'].add(t['user_id'])

        t_key = f"{p_name}::{t_name}"
        if t_key not in task_aggregates:
            task_aggregates[t_key] = {
                'task_name': t_name,
                'project_name': p_name,
                'total_seconds': 0,
                'count': 0,
                'status': status
            }
        task_aggregates[t_key]['total_seconds'] += dur
        task_aggregates[t_key]['count'] += 1

    tasks_total_sec = sum(p['total_seconds'] for p in project_aggregates.values())
    effective_work_seconds = max(total_work_seconds, tasks_total_sec)
    effective_work_hours = round(effective_work_seconds / 3600.0, 2)
    effective_break_hours = round(total_break_seconds / 3600.0, 2)

    by_project = [
        {
            'project_name': p['project_name'],
            'work_hours': round(p['total_seconds'] / 3600.0, 2),
            'work_seconds': p['total_seconds'],
            'work_duration_formatted': format_task_seconds(p['total_seconds']),
            'tasks_count': p['tasks_count'],
            'employee_count': len(p['employees']),
            'percentage': round((p['total_seconds'] / tasks_total_sec * 100), 1) if tasks_total_sec > 0 else (round(100.0 / len(project_aggregates), 1) if project_aggregates else 0)
        }
        for p in sorted(project_aggregates.values(), key=lambda x: x['total_seconds'], reverse=True)
    ]

    by_task = [
        {
            'task_name': tk['task_name'],
            'project_name': tk['project_name'],
            'work_hours': round(tk['total_seconds'] / 3600.0, 2),
            'work_seconds': tk['total_seconds'],
            'work_duration_formatted': format_task_seconds(tk['total_seconds']),
            'count': tk['count'],
            'status': tk['status']
        }
        for tk in sorted(task_aggregates.values(), key=lambda x: x['total_seconds'], reverse=True)[:10]
    ]

    # 6. Employee breakdown (for Admin or multi-user)
    by_employee = []
    if is_admin:
        emp_aggregates = {}
        for u in target_users:
            u_sec = 0
            u_break = 0
            u_tasks = [t for t in scoped_tasks if int(t.get('user_id', -1)) == u.id]
            for d in daily_trend:
                u_day_ev = events_by_user_and_date.get((u.id, date.fromisoformat(d['date'])), [])
                if u_day_ev:
                    m = calculate_day_metrics(u_day_ev, now_utc if d['date'] == today_ist.isoformat() else None)
                    u_sec += m['total_work_seconds']
                    u_break += m['total_break_seconds']

            t_sec = sum(t.get('duration_seconds', 0) for t in u_tasks)
            u_sec = max(u_sec, t_sec)

            if u_sec > 0 or u_tasks:
                emp_code = u.employee_id or f"EMP{u.id:03d}"
                emp_name = f"{u.first_name} {u.last_name}".strip() or u.username
                emp_aggregates[u.id] = {
                    'user_id': u.id,
                    'employee_id': emp_code,
                    'employee_name': emp_name,
                    'department': u.department or 'General',
                    'work_hours': round(u_sec / 3600.0, 2),
                    'work_seconds': u_sec,
                    'work_duration_formatted': format_task_seconds(u_sec),
                    'break_hours': round(u_break / 3600.0, 2),
                    'tasks_count': len(u_tasks),
                    'percentage': round((u_sec / effective_work_seconds * 100), 1) if effective_work_seconds > 0 else 0
                }

        by_employee = sorted(emp_aggregates.values(), key=lambda x: x['work_seconds'], reverse=True)

    # 7. Summary
    total_tasks = len(scoped_tasks)
    active_with_records = len({uid for uid in active_employees_set if uid and uid > 0})
    summary = {
        'total_work_seconds': effective_work_seconds,
        'total_work_hours': effective_work_hours,
        'total_work_formatted': format_task_seconds(effective_work_seconds),
        'total_break_seconds': total_break_seconds,
        'total_break_hours': effective_break_hours,
        'total_break_formatted': format_task_seconds(total_break_seconds),
        'total_tasks': total_tasks,
        'completed_tasks': completed_tasks_count,
        'in_progress_tasks': in_progress_tasks_count,
        'employees_with_records': active_with_records,
        'employee_count': active_with_records if is_admin else 1,
        'total_registered_employees': len(employee_users) if is_admin else 1,
        'active_projects_count': len(project_aggregates),
        'days_worked_count': len(days_worked_set),
        'avg_daily_work_hours': round(effective_work_hours / max(1, len(days_worked_set)), 2) if days_worked_set else 0
    }

    available_employees = []
    if is_admin:
        available_employees = [
            {'id': u.id, 'employee_id': u.employee_id or f"EMP{u.id:03d}", 'name': f"{u.first_name} {u.last_name}".strip() or u.username}
            for u in employee_users
        ]

    all_known_projects = sorted(list(set(
        t.get('project_name', '').strip()
        for t in storage.get_all_tasks()
        if t.get('project_name')
    )))

    return Response({
        'user_role': user.role,
        'is_admin': is_admin,
        'date_range': {
            'start_date': d_start_str,
            'end_date': d_end_str,
            'period': period
        },
        'summary': summary,
        'daily_trend': daily_trend,
        'by_project': by_project,
        'by_employee': by_employee,
        'by_task': by_task,
        'filters': {
            'employees': available_employees,
            'projects': all_known_projects
        }
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_analytics_dashboard_view(request):
    """
    Dedicated Admin Panel Analytics API.
    Strictly forbids non-admin users with 403 Forbidden.
    """
    if request.user.role != 'ADMIN':
        return Response({'error': 'Access denied. Administrator privileges required.'}, status=403)
    return analytics_dashboard_view(getattr(request, '_request', request))




